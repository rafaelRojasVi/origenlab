"""Price-list importer: pure parsing, the plan model and the target guards. No database.

Every workbook and text here is synthetic, built in the test with the real header layout and
invented rows (ACME-1, ACME-2, round invented prices). No real price, discount or margin.
"""
from __future__ import annotations

import json
import sys
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from origenlab_api.v2.catalog import importing
from origenlab_api.v2.catalog.keys import LabdeliveryRefused

_SCRIPTS = Path(__file__).resolve().parents[1] / "scripts" / "catalog"
sys.path.insert(0, str(_SCRIPTS))
import _common  # noqa: E402
import import_price_lists as ipl  # noqa: E402

LOOPBACK_API = "postgresql://origenlab_api:pw@127.0.0.1:5432/origenlab_test_0123abcd"
LOOPBACK_ADMIN = "postgresql://postgres:pw@127.0.0.1:5432/origenlab_test_0123abcd"


# ------------------------------------------------------------------ synthetic inputs

def _ohaus_workbook(path: Path, rows: list[tuple]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2026 List Price"
    ws.append([None])
    ws.append(["ACME Official Dealer Price List"])
    ws.append([None])
    ws.append([None, None, "Effective March 2, 2026"])
    ws.append([None, None, "Last updated February 9, 2026"])
    ws.append(["Category", "Product Family", "Material ID", "Material", "2026 List Price", "2026 Discount %",
               "2026 USD Net Price*", "Release date (2026 releases)", "Notes"])
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    return path


def _adam_workbook(path: Path, rows: list[tuple]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for _ in range(3):
        ws.append([None])
    ws.append(["Información de identificación", None, None, None, "Especificaciones de pesaje", None, None,
               "Precios", None, None])
    ws.append(["ID del producto", "SKU", "UPC", "Nombre", "Capacidad", "Legibilidad", "Tamaño del plato",
               "Precio de lista", "Precio MAP", "Precio del distribuidor"])
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    return path


OHAUS_ROWS = [
    ("Balances", "Acme Family", "10000001", "Compact Scale, ACME-1", 1000, 0.25, 750.0, None, None),
    ("Parts", "Spare Parts", "10000002", "Bubble Level", 20, 0.1, 18.0, None, None),
    ("Parts", "Spare Parts", "10000003", "Calibration Weight, ACME-2", None, None, None, None, None),
]
ADAM_ROWS = [
    ("900001", "ACME 100e", "000000000001", "Balanzas ficticias", "100g", "0.1g", "80mm", "$1,000.00", "$900.00",
     "$600.00"),
    ("900002", "ACME 200i", "000000000002", "Balanzas ficticias", "200g", "0.1g", "80mm", "$2000.00", "", "$1200.50"),
]
LOSER_TEXT = """ACME Messtechnik price list 2024
Pos Description Item no. Price
1 Osmometer ficticio Typ ACME-1 1.000,00 €
2 Thermal printer with power supply AC-2 300,50 €
3 Solucion ficticia 100 mosm/kg (5 ampollas) 9.99.0001 7,00 €
Prices ex factory, net to dealer
"""
ORTOALRESA_TEXT = """TARIFA 2019
CE100 Centrífuga ficticia ACME 1.000,00
RT-20 Rotor ficticio 250,00 €
RE 30 Recambio ficticio 12,50
Condiciones generales
"""


# ------------------------------------------------------------------ parsers

def test_parse_ohaus_xlsx_reads_values_as_decimals(tmp_path: Path) -> None:
    rows = ipl.parse_ohaus_xlsx(_ohaus_workbook(tmp_path / "acme-dealer.xlsx", OHAUS_ROWS))
    assert [r["model_number"] for r in rows] == ["ACME-1", "10000002", "ACME-2"]
    first = rows[0]
    assert first["list_price"] == Decimal("1000") and first["cost"] == Decimal("750")
    assert first["discount_pct"] == Decimal("0.25")
    assert all(isinstance(r[k], Decimal) for r in rows[:2] for k in ("list_price", "cost", "discount_pct"))
    assert (first["currency"], first["as_of"], first["price_kind"], first["is_stale"]) == ("USD", "2026-03-02",
                                                                                           "dealer_net", False)
    assert first["name"] == "Compact Scale, ACME-1" and first["family"] == "Balances / Acme Family"
    assert first["manufacturer"] == first["supplier"] == "OHAUS"
    assert rows[2]["cost"] is None  # no price: kept as a row, dropped by the plan


def test_parse_ohaus_xlsx_needs_its_header(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.active.append(["SKU", "Price"])
    wb.save(tmp_path / "other.xlsx")
    with pytest.raises(ValueError, match="header"):
        ipl.parse_ohaus_xlsx(tmp_path / "other.xlsx")


def test_parse_adam_xlsx_reads_dollar_strings(tmp_path: Path) -> None:
    rows = ipl.parse_adam_xlsx(_adam_workbook(tmp_path / "acme-adam.xlsx", ADAM_ROWS), as_of="2026-08-20")
    assert [r["model_number"] for r in rows] == ["ACME 100e", "ACME 200i"]
    assert rows[0]["list_price"] == Decimal("1000.00") and rows[0]["map_price"] == Decimal("900.00")
    assert rows[0]["cost"] == Decimal("600.00")
    assert rows[1]["map_price"] is None and rows[1]["cost"] == Decimal("1200.50")
    assert (rows[0]["currency"], rows[0]["as_of"], rows[0]["price_kind"]) == ("USD", "2026-08-20", "dealer_net")
    assert rows[0]["family"] == "Balanzas ficticias"


def test_parse_loser_text_comma_decimals() -> None:
    rows = ipl.parse_loser_text(LOSER_TEXT, as_of="2024-01-01")
    assert [(r["model_number"], r["cost"]) for r in rows] == [
        ("Typ ACME-1", Decimal("1000.00")), ("AC-2", Decimal("300.50")), ("9.99.0001", Decimal("7.00"))]
    assert rows[0]["name"] == "Osmometer ficticio"
    assert rows[2]["name"] == "Solucion ficticia 100 mosm/kg (5 ampollas)"
    assert {(r["currency"], r["price_kind"], r["is_stale"]) for r in rows} == {("EUR", "dealer_net", False)}


def test_parse_ortoalresa_text_is_stale() -> None:
    rows = ipl.parse_ortoalresa_text(ORTOALRESA_TEXT, as_of="2019-01-01")
    assert [(r["model_number"], r["list_price"]) for r in rows] == [
        ("CE100", Decimal("1000.00")), ("RT-20", Decimal("250.00")), ("RE 30", Decimal("12.50"))]
    assert all(r["is_stale"] for r in rows)
    assert {(r["currency"], r["price_kind"]) for r in rows} == {("EUR", "list")}
    assert rows[0]["cost"] is None and rows[0]["name"] == "Centrífuga ficticia ACME"


# ------------------------------------------------------------------ Labdelivery (spec S5)

def test_labdelivery_path_refused(tmp_path: Path) -> None:
    folder = tmp_path / "Lab-Delivery 2024"
    folder.mkdir()
    path = _ohaus_workbook(folder / "acme.xlsx", OHAUS_ROWS)
    with pytest.raises(LabdeliveryRefused):
        ipl.build_price_list_plan([("ohaus", path, None)])


def test_labdelivery_first_page_refused(tmp_path: Path) -> None:
    path = tmp_path / "acme.txt"
    path.write_text("LABDELIVERY SpA\n" + LOSER_TEXT, encoding="utf-8")
    with pytest.raises(LabdeliveryRefused):
        ipl.build_price_list_plan([("loser", path, "2024-01-01")])


def test_labdelivery_in_a_written_string_refused() -> None:
    rows = ipl.parse_loser_text("1 Cable de Lab Delivery ACME-9 10,00 €\n", as_of="2024-01-01")
    with pytest.raises(LabdeliveryRefused):
        importing.refuse_labdelivery_in_plan(importing.build_plan("price_lists", [], ipl.plan_items(rows)))


# ------------------------------------------------------------------ the plan

def _plan(tmp_path: Path) -> dict:
    ohaus, adam, loser = tmp_path / "acme-dealer.xlsx", tmp_path / "acme-adam.xlsx", tmp_path / "loser.txt"
    # openpyxl stamps the save time into the file: write each input once, so a second plan reads
    # the same bytes.
    if not ohaus.exists():
        _ohaus_workbook(ohaus, OHAUS_ROWS)
        _adam_workbook(adam, ADAM_ROWS)
        loser.write_text(LOSER_TEXT, encoding="utf-8")
    return ipl.build_price_list_plan([("ohaus", ohaus, None), ("adam", adam, "2026-08-20"),
                                      ("loser", loser, "2024-01-01")])


def test_plan_shape_and_counts(tmp_path: Path) -> None:
    plan = _plan(tmp_path)
    assert plan["importer"] == "price_lists" and plan["version"] == 1
    assert sorted(i["kind"] for i in plan["inputs"]) == ["adam", "loser", "ohaus"]
    # Inputs are planned in file-hash order, so which duplicate wins never depends on flag order.
    assert [i["file_sha256"] for i in plan["inputs"]] == sorted(i["file_sha256"] for i in plan["inputs"])
    assert set(plan["inputs"][0]) == {"path_sha256", "file_sha256", "kind"}
    counts = plan["counts"]
    assert counts["create_org"] == 3  # OHAUS, Adam Equipment, Löser Messtechnik
    assert counts["create_product"] == 2 + 2 + 3
    # ADAM: dealer_net + map for the first row, dealer_net only for the second (no MAP price)
    assert counts["add_observation"] == 2 + 3 + 3
    assert counts["set_terms"] == 1
    assert counts["rows_without_price"] == 1
    actions = [i["action"] for i in plan["items"]]
    assert actions == sorted(actions, key=importing.ACTION_ORDER.index)
    [terms] = [i for i in plan["items"] if i["action"] == "set_terms"]
    assert terms["fields"]["map_enforced"] is True and terms["fields"]["currency"] == "USD"
    kinds = sorted(i["fields"]["price_kind"] for i in plan["items"]
                   if i["action"] == "add_observation" and i["fields"]["product"].endswith("ACME100E"))
    assert kinds == ["dealer_net", "map"]


def test_plan_never_contains_local_paths(tmp_path: Path) -> None:
    plan = _plan(tmp_path)
    text = importing.plan_bytes(plan).decode()
    assert str(tmp_path) not in text and "acme-dealer.xlsx" not in text and "loser.txt" not in text


def test_plan_is_deterministic(tmp_path: Path) -> None:
    assert importing.plan_sha256(_plan(tmp_path)) == importing.plan_sha256(_plan(tmp_path))


def test_money_is_serialised_as_strings(tmp_path: Path) -> None:
    plan = json.loads(importing.plan_bytes(_plan(tmp_path)))
    obs = next(i for i in plan["items"] if i["action"] == "add_observation")
    assert isinstance(obs["fields"]["price"], str)


def test_conflicting_duplicate_keeps_first_and_counts() -> None:
    a = {"action": "create_org", "key": "org:acme", "fields": {"name": "ACME"}, "source": {"input": "a", "ref": "1"}}
    same = {**a, "source": {"input": "b", "ref": "2"}}
    b = {"action": "create_org", "key": "org:acme", "fields": {"name": "Acme"}, "source": {"input": "c", "ref": "3"}}
    plan = importing.build_plan("t", [], [a, same, b])
    assert plan["counts"]["create_org"] == 1
    assert plan["counts"]["duplicates_merged"] == 1 and plan["counts"]["conflicting_duplicates"] == 1
    assert plan["conflicts"] == [{"key": "org:acme", "kept": {"input": "a", "ref": "1"},
                                  "dropped": [{"input": "c", "ref": "3"}]}]
    assert plan["items"][0]["fields"]["name"] == "ACME"


def test_duplicate_winner_does_not_depend_on_flag_order(tmp_path: Path) -> None:
    first, second = tmp_path / "first.txt", tmp_path / "second.txt"
    first.write_text("1 Osmometer ficticio ACME-1 100,00 €\n", encoding="utf-8")
    second.write_text("1 Osmometer ficticio ACME-1 200,00 €\n", encoding="utf-8")
    one = ipl.build_price_list_plan([("loser", first, "2024-01-01"), ("loser", second, "2024-01-01")])
    two = ipl.build_price_list_plan([("loser", second, "2024-01-01"), ("loser", first, "2024-01-01")])
    assert importing.plan_sha256(one) == importing.plan_sha256(two)
    [conflict] = one["conflicts"]
    winner = min(importing.input_record(p, "loser")["file_sha256"] for p in (first, second))
    assert conflict["key"].startswith("obs:") and conflict["kept"] == {"input": winner, "ref": "1"}


def test_ohaus_rows_carry_their_material_id(tmp_path: Path) -> None:
    rows = ipl.parse_ohaus_xlsx(_ohaus_workbook(tmp_path / "acme.xlsx", OHAUS_ROWS))
    assert rows[0]["source_ref"] == "10000001"


@pytest.mark.parametrize("discount", [1, 1.0, 25, -0.1])
def test_ohaus_discount_must_be_a_fraction_below_one(tmp_path: Path, discount: float) -> None:
    row = ("Balances", "Acme Family", "10000009", "Compact Scale, ACME-9", 1000, discount, 750.0, None, None)
    with pytest.raises(ValueError, match="10000009"):
        ipl.parse_ohaus_xlsx(_ohaus_workbook(tmp_path / "acme.xlsx", [row]))


def test_load_plan_refuses_a_mismatched_sha(tmp_path: Path) -> None:
    data = importing.plan_bytes(_plan(tmp_path))
    assert importing.load_plan(data, importing.sha256_bytes(data))["importer"] == "price_lists"
    with pytest.raises(importing.PlanRefused):
        importing.load_plan(data, "0" * 64)


def test_compare_reports_present_missing_different() -> None:
    planned = {"price": "10.00", "currency": "USD", "is_stale": False}
    assert importing.compare(planned, None, ("price",))[0] == "missing"
    assert importing.compare(planned, {"price": Decimal("10"), "currency": "USD"}, ("price", "currency"))[0] == "present"
    status, diffs = importing.compare(planned, {"price": Decimal("11"), "currency": "USD"}, ("price", "currency"))
    assert status == "different" and diffs == ["price"]


# ------------------------------------------------------------------ the CLI guards (no database reached)

def _write_plan(tmp_path: Path) -> tuple[Path, str]:
    data = importing.plan_bytes(_plan(tmp_path))
    out = tmp_path / "plan.json"
    out.write_bytes(data)
    return out, importing.sha256_bytes(data)


def _apply_argv(tmp_path: Path, plan: Path, sha: str, *, target: str = LOOPBACK_API,
                admin: str = LOOPBACK_ADMIN, extra: tuple[str, ...] = ()) -> list[str]:
    return ["apply", "--target-dsn", target, "--admin-dsn", admin, "--operator-email", "op@example.test",
            "--plan", str(plan), "--plan-sha256", sha, "--out", str(tmp_path / "out"), *extra]


def test_hosted_dsn_refused(tmp_path: Path) -> None:
    plan, sha = _write_plan(tmp_path)
    hosted = "postgresql://x@db.abc.supabase.co/postgres"
    assert ipl.main(_apply_argv(tmp_path, plan, sha, target=hosted)) == _common.EXIT_REFUSED
    assert ipl.main(_apply_argv(tmp_path, plan, sha, admin=hosted)) == _common.EXIT_REFUSED
    assert ipl.main(["verify", "--target-dsn", hosted, "--plan", str(plan), "--plan-sha256", sha,
                     "--out", str(tmp_path / "v")]) == _common.EXIT_REFUSED


def test_host_name_refused_even_localhost(tmp_path: Path) -> None:
    plan, sha = _write_plan(tmp_path)
    named = "postgresql://origenlab_api:pw@localhost:5432/origenlab_test_0123abcd"
    assert ipl.main(_apply_argv(tmp_path, plan, sha, target=named)) == _common.EXIT_REFUSED


def test_plan_sha_mismatch_refused(tmp_path: Path) -> None:
    plan, _ = _write_plan(tmp_path)
    assert ipl.main(_apply_argv(tmp_path, plan, "f" * 64)) == _common.EXIT_REFUSED
    assert ipl.main(["verify", "--target-dsn", LOOPBACK_API, "--plan", str(plan), "--plan-sha256", "f" * 64,
                     "--out", str(tmp_path / "v")]) == _common.EXIT_REFUSED


def test_non_disposable_target_refused_without_cleanroom_flag(tmp_path: Path) -> None:
    plan, sha = _write_plan(tmp_path)
    clean_api = LOOPBACK_API.replace("origenlab_test_0123abcd", "origenlab_clean")
    clean_admin = LOOPBACK_ADMIN.replace("origenlab_test_0123abcd", "origenlab_clean")
    assert ipl.main(_apply_argv(tmp_path, plan, sha, target=clean_api, admin=clean_admin)) == _common.EXIT_REFUSED
    other = LOOPBACK_API.replace("origenlab_test_0123abcd", "origenlab_dev")
    assert ipl.main(_apply_argv(tmp_path, plan, sha, target=other,
                                admin=other, extra=("--allow-cleanroom-production",))) == _common.EXIT_REFUSED


def test_target_and_admin_must_name_one_database(tmp_path: Path) -> None:
    plan, sha = _write_plan(tmp_path)
    admin = LOOPBACK_ADMIN.replace("0123abcd", "89abcdef")
    assert ipl.main(_apply_argv(tmp_path, plan, sha, admin=admin)) == _common.EXIT_REFUSED


def test_out_dir_inside_the_repository_refused(tmp_path: Path) -> None:
    inside = _common.REPO / "apps" / "api" / ".catalog-import-out-test"
    ohaus = _ohaus_workbook(tmp_path / "acme.xlsx", OHAUS_ROWS)
    assert ipl.main(["plan", "--ohaus", str(ohaus), "--out", str(inside)]) == _common.EXIT_REFUSED
    assert not inside.exists()


def test_out_dir_inside_any_git_checkout_refused(tmp_path: Path) -> None:
    (tmp_path / "other-repo" / ".git").mkdir(parents=True)
    inside = tmp_path / "other-repo" / "reports"
    ohaus = _ohaus_workbook(tmp_path / "acme.xlsx", OHAUS_ROWS)
    assert ipl.main(["plan", "--ohaus", str(ohaus), "--out", str(inside)]) == _common.EXIT_REFUSED
    assert not inside.exists()


def test_messages_never_print_local_paths(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    missing = tmp_path / "private" / "absent-list.xlsx"
    assert ipl.main(["plan", "--ohaus", str(missing), "--out", str(tmp_path / "o")]) == _common.EXIT_REFUSED
    printed = capsys.readouterr()
    assert str(tmp_path) not in printed.out + printed.err and "absent-list.xlsx" in printed.err
    ohaus = _ohaus_workbook(tmp_path / "acme.xlsx", OHAUS_ROWS)
    assert ipl.main(["plan", "--ohaus", str(ohaus), "--out", str(tmp_path / "o2")]) == 0
    assert ipl.main(["plan", "--ohaus", str(ohaus), "--out", str(tmp_path / "o2")]) == _common.EXIT_REFUSED
    plan = tmp_path / "o2" / "plan.json"
    assert ipl.main(["verify", "--target-dsn", LOOPBACK_API, "--plan", str(tmp_path / "nope.json"),
                     "--plan-sha256", "0" * 64, "--out", str(tmp_path / "v")]) == _common.EXIT_REFUSED
    printed = capsys.readouterr()
    assert str(tmp_path) not in printed.out + printed.err and "plan.json" in printed.out
    assert plan.exists()


def test_plan_command_writes_plan_and_prints_its_sha(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    ohaus = _ohaus_workbook(tmp_path / "acme.xlsx", OHAUS_ROWS)
    out = tmp_path / "out"
    assert ipl.main(["plan", "--ohaus", str(ohaus), "--out", str(out)]) == 0
    data = (out / "plan.json").read_bytes()
    assert importing.sha256_bytes(data) in capsys.readouterr().out
    assert str(tmp_path) not in data.decode()
    # A second plan into the same directory never overwrites the first.
    assert ipl.main(["plan", "--ohaus", str(ohaus), "--out", str(out)]) == _common.EXIT_REFUSED


def test_rollback_needs_its_confirmation_flag(tmp_path: Path) -> None:
    plan, sha = _write_plan(tmp_path)
    argv = ["rollback", "--target-dsn", LOOPBACK_API, "--admin-dsn", LOOPBACK_ADMIN, "--plan", str(plan),
            "--plan-sha256", sha, "--out", str(tmp_path / "r")]
    assert ipl.main(argv) == _common.EXIT_REFUSED
