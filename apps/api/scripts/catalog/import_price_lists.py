#!/usr/bin/env python3
"""Load supplier price lists into the catalog: products and their supplier cost observations.

    uv run python scripts/catalog/import_price_lists.py plan --out ~/data/…/catalog-import-<ts> \\
        --ohaus <dealer price list.xlsx> [--ohaus …] --adam <price list.xlsx> --as-of adam=YYYY-MM-DD \\
        --loser-text <price list text.txt> --as-of loser=YYYY-MM-DD \\
        --ortoalresa-text <tariff text.txt> --as-of ortoalresa=YYYY-MM-DD
    uv run python scripts/catalog/import_price_lists.py apply  --target-dsn … --admin-dsn … \\
        --operator-email … --plan …/plan.json --plan-sha256 <sha> --out …
    uv run python scripts/catalog/import_price_lists.py verify   --target-dsn … --plan … --plan-sha256 … --out …
    uv run python scripts/catalog/import_price_lists.py rollback --target-dsn … --admin-dsn … --plan … \\
        --plan-sha256 … --out … --confirm-delete-loaded-rows

Targets, exit codes and the manifest: `_common.py`. PDF price lists are read as text already
extracted (`pdftotext -layout`), so no PDF library is needed here.

What each list becomes (one organization per maker, which is also the supplier):

* OHAUS dealer workbook (header "Material ID", "Material", "List Price", "Discount", "Net Price";
  date from its "Effective <Month D, YYYY>" line, or `--as-of ohaus=`): one `dealer_net`
  observation per row (price = net, with list price and discount).
* ADAM workbook (header "SKU", "Nombre", "Precio de lista", "Precio MAP", "Precio del
  distribuidor"): a `dealer_net` observation (distributor price, with list price) and a `map`
  observation (MAP price) per row, and supplier terms with `map_enforced = true` if the supplier
  has none.
* Löser text (`<pos> <description> <item no> <price>` with comma decimals, EUR): `dealer_net`.
* Ortoalresa text (`<CE|RT|RE code> <description> <price>`, EUR): `list`, marked `is_stale`.

Rows without a price are counted (`rows_without_price`) and skipped. Every input passes the
Labdelivery refusal on its path and first page, and every string the plan would write is checked.
"""
from __future__ import annotations

import argparse
import re
import sys
import uuid
from collections import Counter
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, TypedDict

sys.path.insert(0, str(Path(__file__).resolve().parent))

from origenlab_api.v2.catalog import importing  # noqa: E402
from origenlab_api.v2.catalog.keys import model_key, refuse_labdelivery  # noqa: E402

import _common  # noqa: E402

IMPORTER = "price_lists"
KINDS = ("ohaus", "adam", "loser", "ortoalresa")
DEFAULT_ORG = {"ohaus": "OHAUS", "adam": "Adam Equipment", "loser": "Löser Messtechnik", "ortoalresa": "Ortoalresa"}
ROUTES = ("import_courier", "import_freight", "domestic")
SOURCE_DOCUMENT = "price list"
_SIX = Decimal("0.000001")
_EFFECTIVE = re.compile(r"effective\s+([A-Za-z]+\s+\d{1,2},\s*\d{4})", re.IGNORECASE)
_CURRENCY = re.compile(r"\b(USD|EUR|CLP)\b")
_EURO = r"\d{1,3}(?:\.\d{3})*,\d{2}"
_LOSER_LINE = re.compile(rf"^\s*(?P<pos>\d+)\s+(?P<desc>.+?)\s+(?P<item>Typ\s+\S+|\S+)\s+(?:€\s*)?(?P<price>{_EURO})\s*(?:€|EUR)?\s*$")
_ORTOALRESA_LINE = re.compile(
    r"^\s*(?P<code>(?:CE|RT|RE)[\s-]?\d[\w.\-/]*)\s+(?P<desc>.+?)\s+"
    r"(?P<price>\d{1,3}(?:\.\d{3})+(?:,\d{1,2})?|\d+(?:,\d{1,2})?)\s*(?:€|EUR)?\s*$")


class Row(TypedDict):
    manufacturer: str
    supplier: str
    model_number: str
    name: str | None
    family: str | None
    list_price: Decimal | None
    cost: Decimal | None
    map_price: Decimal | None
    discount_pct: Decimal | None
    currency: str
    as_of: str
    is_stale: bool
    price_kind: str
    #: The row's own reference in its list (Material ID, product ID, position, code): plan provenance only.
    source_ref: str | None


# ------------------------------------------------------------------ values

def _money(value: Any) -> Decimal | None:
    """A cell or text amount: numbers as printed, "$1,234.50" strings; quantised as the columns store them."""
    if isinstance(value, str):
        value = value.replace("$", "").replace(",", "").replace("€", "").strip()
    number = importing.to_decimal(value)
    return None if number is None else number.quantize(_SIX)


def _euro(text: str) -> Decimal:
    """"1.234,56" → Decimal("1234.56")."""
    return Decimal(text.replace(".", "").replace(",", ".")).quantize(_SIX)


def _fraction(value: Any, ref: str) -> Decimal | None:
    """The OHAUS "Discount %" cell: a fraction of the list price in [0, 1) (0.25 is 25 %).

    That is how the dealer workbook stores it (a percent-formatted cell). Anything else — 1 (which
    would be either 100 % or a mis-typed 1 %), 25, a negative — is refused with the row's Material
    ID rather than guessed; `catalog.supplier_product.discount_pct` holds [0, 1) as well.
    """
    number = _money(value)
    if number is None:
        return None
    if not Decimal(0) <= number < 1:
        raise ValueError(f"Material ID {ref}: the discount is not a fraction in [0, 1)")
    return number


def _as_of(value: str | None, what: str) -> str:
    if not value:
        raise ValueError(f"{what}: no date for the price list; give --as-of {what}=YYYY-MM-DD")
    return date.fromisoformat(value).isoformat()


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


# ------------------------------------------------------------------ workbooks

def _sheet_rows(path: Path) -> list[tuple[Any, ...]]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [tuple(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    finally:
        wb.close()


def _header(rows: list[tuple[Any, ...]], wanted: dict[str, Any]) -> tuple[int, dict[str, int]]:
    """The first row that has every wanted column; `wanted` maps a name to a predicate on the cell text."""
    for n, row in enumerate(rows):
        cells = [_text(c) for c in row]
        found: dict[str, int] = {}
        for name, test in wanted.items():
            idx = next((i for i, c in enumerate(cells) if test(c)), None)
            if idx is not None:
                found[name] = idx
        if len(found) == len(wanted):
            return n, found
    raise ValueError(f"no header row with {', '.join(wanted)}")


def xlsx_first_page(path: Path) -> str:
    """The workbook's first sheet title and first rows, as text: what the Labdelivery check reads."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        lines = [" ".join(wb.sheetnames)]
        for n, row in enumerate(ws.iter_rows(values_only=True)):
            if n >= 20:
                break
            lines.append(" ".join(_text(c) for c in row if c is not None))
        return "\n".join(lines)
    finally:
        wb.close()


def parse_ohaus_xlsx(path: Path, *, as_of: str | None = None, org: str = DEFAULT_ORG["ohaus"]) -> list[Row]:
    rows = _sheet_rows(Path(path))
    header_at, col = _header(rows, {
        "material_id": lambda c: c.lower() == "material id", "material": lambda c: c.lower() == "material",
        "list": lambda c: "list price" in c.lower(), "discount": lambda c: "discount" in c.lower(),
        "net": lambda c: "net price" in c.lower()})
    preamble = " ".join(_text(c) for r in rows[:header_at] for c in r if c is not None)
    effective = _EFFECTIVE.search(preamble)
    if as_of is None and effective:
        as_of = datetime.strptime(re.sub(r"\s+", " ", effective.group(1)), "%B %d, %Y").date().isoformat()
    as_of = _as_of(as_of, "ohaus")
    net_header = _text(rows[header_at][col["net"]])
    currency = (_CURRENCY.search(net_header) or _CURRENCY.search("USD")).group(1)
    category = 0 if _text(rows[header_at][0]).lower() == "category" else None
    family_idx = next((i for i, c in enumerate(rows[header_at]) if "family" in _text(c).lower()), None)

    out: list[Row] = []
    for row in rows[header_at + 1:]:
        cell = lambda i: row[i] if i is not None and i < len(row) else None  # noqa: E731
        material_id, material = _text(cell(col["material_id"])), _text(cell(col["material"]))
        if not material_id or not material:
            continue
        model = material.rsplit(",", 1)[1].strip() if "," in material else material_id
        family = " / ".join(x for x in (_text(cell(category)), _text(cell(family_idx))) if x) or None
        out.append(Row(manufacturer=org, supplier=org, model_number=model, name=material, family=family,
                       list_price=_money(cell(col["list"])), cost=_money(cell(col["net"])), map_price=None,
                       discount_pct=_fraction(cell(col["discount"]), material_id), currency=currency, as_of=as_of,
                       is_stale=False, price_kind="dealer_net", source_ref=material_id))
    return out


def parse_adam_xlsx(path: Path, *, as_of: str | None = None, org: str = DEFAULT_ORG["adam"]) -> list[Row]:
    rows = _sheet_rows(Path(path))
    header_at, col = _header(rows, {
        "sku": lambda c: c.lower() == "sku", "name": lambda c: c.lower() == "nombre",
        "list": lambda c: c.lower() == "precio de lista", "map": lambda c: c.lower() == "precio map",
        "dealer": lambda c: c.lower() == "precio del distribuidor"})
    as_of = _as_of(as_of, "adam")
    id_idx = next((i for i, c in enumerate(rows[header_at]) if _text(c).lower() == "id del producto"), None)
    out: list[Row] = []
    for row in rows[header_at + 1:]:
        cell = lambda i: row[i] if i < len(row) else None  # noqa: E731
        sku = _text(cell(col["sku"]))
        if not sku:
            continue
        family = _text(cell(col["name"])) or None
        out.append(Row(manufacturer=org, supplier=org, model_number=sku, name=f"{family or ''} {sku}".strip(),
                       family=family, list_price=_money(cell(col["list"])), cost=_money(cell(col["dealer"])),
                       map_price=_money(cell(col["map"])), discount_pct=None, currency="USD", as_of=as_of,
                       is_stale=False, price_kind="dealer_net", source_ref=(_text(cell(id_idx)) if id_idx is not None else "") or sku))
    return out


# ------------------------------------------------------------------ extracted PDF text

def parse_loser_text(text: str, *, as_of: str | None = None, org: str = DEFAULT_ORG["loser"]) -> list[Row]:
    as_of = _as_of(as_of, "loser")
    out: list[Row] = []
    for line in text.splitlines():
        m = _LOSER_LINE.match(line)
        if not m:
            continue
        out.append(Row(manufacturer=org, supplier=org, model_number=re.sub(r"\s+", " ", m["item"]),
                       name=m["desc"].strip(), family=None, list_price=None, cost=_euro(m["price"]), map_price=None,
                       discount_pct=None, currency="EUR", as_of=as_of, is_stale=False, price_kind="dealer_net",
                       source_ref=m["pos"]))
    return out


def parse_ortoalresa_text(text: str, *, as_of: str | None = None, org: str = DEFAULT_ORG["ortoalresa"]) -> list[Row]:
    as_of = _as_of(as_of, "ortoalresa")
    out: list[Row] = []
    for line in text.splitlines():
        m = _ORTOALRESA_LINE.match(line)
        if not m:
            continue
        out.append(Row(manufacturer=org, supplier=org, model_number=re.sub(r"\s+", " ", m["code"]),
                       name=m["desc"].strip(), family=None, list_price=_euro(m["price"]), cost=None, map_price=None,
                       discount_pct=None, currency="EUR", as_of=as_of, is_stale=True, price_kind="list",
                       source_ref=re.sub(r"\s+", " ", m["code"])))
    return out


def text_first_page(text: str) -> str:
    return text.split("\f", 1)[0][:20000]


# ------------------------------------------------------------------ the plan

def _org_key(name: str) -> str:
    return "org:" + importing.norm_name(name)


def plan_items(rows: list[Row], *, org_ids: dict[str, str] | None = None,
               skipped: Counter[str] | None = None, input_sha256: str | None = None) -> list[importing.PlanItem]:
    """Organizations, products and observations for parsed rows; rows without a price are counted, not planned."""
    ids = {importing.norm_name(k): v for k, v in (org_ids or {}).items()}
    skipped = skipped if skipped is not None else Counter()
    items: list[importing.PlanItem] = []
    planned_orgs: set[str] = set()

    def org(name: str) -> str:
        key = _org_key(name)
        if key not in planned_orgs:
            planned_orgs.add(key)
            fields: dict[str, Any] = {"name": name, "kind": "supplier"}
            if importing.norm_name(name) in ids:
                fields["organization_id"] = ids[importing.norm_name(name)]
            items.append({"action": "create_org", "key": key, "fields": fields})
        return key

    for r in rows:
        price = r["list_price"] if r["price_kind"] == "list" else r["cost"]
        key = model_key(r["model_number"])
        if price is None:
            skipped["rows_without_price"] += 1
            continue
        if not key:
            skipped["rows_without_model"] += 1
            continue
        maker, supplier = org(r["manufacturer"]), org(r["supplier"])
        source = {"input": input_sha256, "ref": r.get("source_ref")}
        product = f"product:{importing.norm_name(r['manufacturer'])}|{key}"
        items.append({"action": "create_product", "key": product, "source": source,
                      "fields": {"manufacturer": maker, "model_number": r["model_number"], "name": r["name"],
                                 "description": r["family"]}})
        observations = [(r["price_kind"], price, r["discount_pct"])]
        if r["map_price"] is not None:
            observations.append(("map", r["map_price"], None))
        for kind, amount, discount in observations:
            items.append({"action": "add_observation",
                          "key": f"obs:{importing.norm_name(r['supplier'])}|{key}|{r['as_of']}|{kind}",
                          "source": source,
                          "fields": {"supplier": supplier, "product": product, "as_of": r["as_of"],
                                     "price_kind": kind, "price": amount, "currency": r["currency"],
                                     "list_price": r["list_price"], "discount_pct": discount,
                                     "is_stale": r["is_stale"], "source_document": SOURCE_DOCUMENT}})
    return items


def _read_input(kind: str, path: Path) -> tuple[str, str | None]:
    """(first page, full text) — workbooks have no full text here; they are parsed from the file."""
    if kind in ("ohaus", "adam"):
        return xlsx_first_page(path), None
    text = path.read_text(encoding="utf-8")
    return text_first_page(text), text


def build_price_list_plan(inputs: list[tuple[str, Path, str | None]], *, org_names: dict[str, str] | None = None,
                          org_ids: dict[str, str] | None = None,
                          adam_route: str = "import_courier") -> importing.Plan:
    """`inputs` is (kind, path, as_of) per file; every file is refused if it is of Labdelivery origin."""
    names = {**DEFAULT_ORG, **(org_names or {})}
    records, items = [], []
    skipped: Counter[str] = Counter()
    # Planned in file-hash order: which of two conflicting rows wins never depends on flag order.
    hashed = []
    for kind, path, as_of in inputs:
        if kind not in KINDS:
            raise ValueError(f"unknown price list kind {kind!r}")
        hashed.append((importing.input_record(Path(path), kind), kind, Path(path), as_of))
    hashed.sort(key=lambda h: (h[0]["file_sha256"], h[1]))
    for record, kind, path, as_of in hashed:
        first_page, text = _read_input(kind, path)
        refuse_labdelivery(str(path.resolve()), path.name, first_page)
        if kind == "ohaus":
            rows = parse_ohaus_xlsx(path, as_of=as_of, org=names[kind])
        elif kind == "adam":
            rows = parse_adam_xlsx(path, as_of=as_of, org=names[kind])
        elif kind == "loser":
            rows = parse_loser_text(text or "", as_of=as_of, org=names[kind])
        else:
            rows = parse_ortoalresa_text(text or "", as_of=as_of, org=names[kind])
        records.append(record)
        items += plan_items(rows, org_ids=org_ids, skipped=skipped, input_sha256=record["file_sha256"])
        if kind == "adam" and rows:
            items.append({"action": "set_terms", "key": f"terms:{importing.norm_name(names[kind])}",
                          "fields": {"supplier": _org_key(names[kind]), "currency": "USD", "route": adam_route,
                                     "map_enforced": True}})
    skipped.setdefault("rows_without_price", 0)
    return importing.build_plan(IMPORTER, records, items, extra_counts=skipped)


# ------------------------------------------------------------------ the CLI

def _pairs(values: list[str], what: str) -> dict[str, str]:
    out: dict[str, str] = {}
    for value in values or []:
        name, sep, rest = value.partition("=")
        if not sep or not name.strip() or not rest.strip():
            raise ValueError(f"{what} {value!r}: want NAME=VALUE")
        out[name.strip()] = rest.strip()
    return out


def _add_plan_arguments(p: argparse.ArgumentParser) -> None:
    p.add_argument("--ohaus", type=Path, action="append", default=[], help="OHAUS dealer price list (.xlsx)")
    p.add_argument("--adam", type=Path, action="append", default=[], help="ADAM price list (.xlsx)")
    p.add_argument("--loser-text", type=Path, action="append", default=[], help="Löser price list, as text")
    p.add_argument("--ortoalresa-text", type=Path, action="append", default=[], help="Ortoalresa tariff, as text")
    p.add_argument("--as-of", action="append", default=[], metavar="KIND=YYYY-MM-DD",
                   help="the date a list is valid from (OHAUS reads its own 'Effective' line)")
    p.add_argument("--org-name", action="append", default=[], metavar="KIND=NAME",
                   help="the organization a kind of list belongs to (default: " + ", ".join(
                       f"{k}={v}" for k, v in DEFAULT_ORG.items()) + ")")
    p.add_argument("--org-id", action="append", default=[], metavar="NAME=UUID",
                   help="use this existing organization for NAME instead of matching by name")
    p.add_argument("--adam-route", choices=ROUTES, default="import_courier",
                   help="import route of the ADAM supplier terms, if the supplier has none")


def _make_plan(args: argparse.Namespace) -> importing.Plan:
    as_of = _pairs(args.as_of, "--as-of")
    org_names = _pairs(args.org_name, "--org-name")
    org_ids = _pairs(args.org_id, "--org-id")
    for name, value in org_ids.items():
        try:
            org_ids[name] = str(uuid.UUID(value))
        except ValueError:
            raise ValueError(f"--org-id for {name!r} is not a UUID") from None
    unknown = sorted((set(as_of) | set(org_names)) - set(KINDS))
    if unknown:
        raise ValueError(f"unknown kind(s) {unknown}; kinds are {', '.join(KINDS)}")
    inputs = [(kind, path, as_of.get(kind)) for kind, paths in (
        ("ohaus", args.ohaus), ("adam", args.adam), ("loser", args.loser_text), ("ortoalresa", args.ortoalresa_text))
        for path in paths]
    if not inputs:
        raise ValueError("no price list given")
    return build_price_list_plan(inputs, org_names=org_names, org_ids=org_ids, adam_route=args.adam_route)


def main(argv: list[str] | None = None) -> int:
    return _common.main(argv, importer=IMPORTER, description=__doc__ or "", add_plan_arguments=_add_plan_arguments,
                        make_plan=_make_plan)


if __name__ == "__main__":
    raise SystemExit(main())
