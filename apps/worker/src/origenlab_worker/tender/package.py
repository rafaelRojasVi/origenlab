"""The pure pieces of a tender's `fill` phase (spec 2026-10-19 §2.3, §6.2; plan Task 12): the
economic offer read from the approved Cotizador revision, which checklist lines the bidder's
standard documents and the spec sheets answer, each line's estado, and the package — one zip whose
members keep the bases' own file names, numbered in the checklist's order.

Nothing here touches the database, Storage, Drive or a model; `job.py` does, and passes the bytes.
"""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Callable, Iterable
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal
from typing import Any, BinaryIO, Final
from zoneinfo import ZoneInfo

#: Every output kind a fill writes (the dashboard labels them).
OUTPUT_KINDS: Final = ("formulario", "formulario_generado", "oferta_economica", "ficha_tecnica",
                       "documento_oferente", "paquete")
CHECKLIST_STATES: Final = ("listo", "falta", "firmar")
#: The whole package, summed before it is written: a bigger one fails `package_too_large`. Below
#: hosted Storage's default per-object limit (50 MB), which the zip is stored under.
MAX_PACKAGE_BYTES: int = 45 * 1024 * 1024
#: One spec sheet downloaded from Drive.
MAX_ATTACHMENT_BYTES: int = 20 * 1024 * 1024
MAX_NAME: Final = 120

SANTIAGO: Final = ZoneInfo("America/Santiago")
CONTENT_TYPES: Final = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "jpg": "image/jpeg",
    "png": "image/png",
    "zip": "application/zip",
}
EXTENSIONS: Final = {v: k for k, v in CONTENT_TYPES.items() if k != "zip"}

_OFFER = re.compile(r"oferta\s+econ[oó]mica|formulario\s+econ", re.IGNORECASE)
_DATASHEET = re.compile(r"fichas?\s+t[eé]cnicas?|cat[aá]logos?|folletos?", re.IGNORECASE)
_BIDDER_KINDS: Final = (
    ("cedula", re.compile(r"c[eé]dula|carnet\s+de\s+identidad", re.IGNORECASE)),
    ("certificado_inscripcion", re.compile(
        r"inscripci[oó]n|chile\s*proveedores|registro\s+de\s+proveedores|certificado\s+de\s+habilidad", re.IGNORECASE)),
    ("declaracion_jurada", re.compile(r"declaraci[oó]n(es)?\s+jurada", re.IGNORECASE)),
    ("experiencia", re.compile(
        r"experiencia|(listado|cartera|n[oó]mina)\s+de\s+clientes|referencias\s+comerciales|cartas?\s+de\s+recomendaci",
        re.IGNORECASE)),
)
_UNSAFE = re.compile(r'[<>:"|?*]')
_CONTROL = re.compile(r"[\x00-\x1f\x7f]")


class PackageTooLarge(Exception):
    """The package would exceed `MAX_PACKAGE_BYTES`."""


# ── which checklist line is what ────────────────────────────────────────────────────────────


def is_offer_entry(entry: dict[str, Any]) -> bool:
    """The line of the economic offer form: a `formulario` named as the economic offer."""
    return entry.get("tipo") == "formulario" and bool(_OFFER.search(str(entry.get("documento") or "")))


def is_datasheet_entry(entry: dict[str, Any]) -> bool:
    """A line asking for the offered products' spec sheets (not a form of the buyer's to fill)."""
    return not entry.get("archivo_formulario_id") and bool(_DATASHEET.search(str(entry.get("documento") or "")))


def bidder_kinds_for(entry: dict[str, Any]) -> set[str]:
    """The `crm.bidder_document` kinds that answer a checklist line. A line whose form is one of
    the anexos is filled instead, never answered with a standard document."""
    if entry.get("archivo_formulario_id"):
        return set()
    text = str(entry.get("documento") or "")
    return {kind for kind, pattern in _BIDDER_KINDS if pattern.search(text)}


def checklist_status(entry: dict[str, Any], outputs: list[dict[str, Any]], bidder_documents: list[Any]) -> str:
    """`falta` when nothing answers the line yet (no output and no bidder document) or an output
    still carries a «[completar: …]»; `firmar` when the bases ask it signed; else `listo`. A line
    with nothing to sign yet is `falta`, not `firmar`: the person must first get the document."""
    if not outputs and not bidder_documents:
        return "falta"
    if any(int(o.get("completar_count") or 0) > 0 for o in outputs):
        return "falta"
    if entry.get("requiere_firma"):
        return "firmar"
    return "listo"


# ── the economic offer ──────────────────────────────────────────────────────────────────────


def _money(value: Decimal, decimals: int) -> int | str:
    q = Decimal(value).quantize(Decimal(1) if decimals == 0 else Decimal("0.01"), ROUND_HALF_UP)
    return int(q) if decimals == 0 else f"{q:.2f}"


def _quantity(value: Any) -> int | str:
    d = Decimal(value).normalize()
    return int(d) if d == d.to_integral_value() else format(d, "f")


def _date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return (value.astimezone(SANTIAGO) if value.tzinfo else value).date().isoformat()
    return None


def offer_payload(revision: dict[str, Any]) -> dict[str, Any]:
    """The offer as the forms print it, from the approved Cotizador revision: number, date, and per
    line and in total the neto, IVA (the revision's `tax_rate`) and total. CLP, rounded half up per
    line as the quote stores it; the totals are the revision's own. Logistics lines are the quote's
    internal costs (spread into the item prices) and are not printed."""
    decimals = int(revision.get("price_decimals") or 0)
    rate = Decimal(revision.get("tax_rate") if revision.get("tax_rate") is not None else "0.19")
    lineas: list[dict[str, Any]] = []
    net_sum = Decimal(0)
    for line in revision.get("lines") or []:
        if line.get("line_kind") == "logistics" or line.get("line_total_qc") is None:
            continue
        neto = Decimal(line["line_total_qc"])
        quantity = Decimal(line.get("quantity") or 1)
        unit = Decimal(line["unit_price_qc"]) if line.get("unit_price_qc") is not None else neto / quantity
        iva = (neto * rate).quantize(Decimal(1) if decimals == 0 else Decimal("0.01"), ROUND_HALF_UP)
        net_sum += neto
        lineas.append({"descripcion": str(line.get("description") or ""), "cantidad": _quantity(quantity),
                       "precio_unitario_neto": _money(unit, decimals), "neto": _money(neto, decimals),
                       "iva": _money(iva, decimals), "total": _money(neto + iva, decimals)})
    neto = Decimal(revision["tax_base"]) if revision.get("tax_base") is not None else net_sum
    iva = Decimal(revision["tax_total"]) if revision.get("tax_total") is not None else neto * rate
    total = Decimal(revision["grand_total"]) if revision.get("grand_total") is not None else neto + _q(iva, decimals)
    return {"numero": revision.get("quote_number"),
            "fecha": _date(revision.get("approved_at")) or _date(revision.get("created_at")),
            "lineas": lineas, "neto": _money(neto, decimals), "iva": _money(iva, decimals),
            "total": _money(total, decimals)}


def _q(value: Decimal, decimals: int) -> Decimal:
    return value.quantize(Decimal(1) if decimals == 0 else Decimal("0.01"), ROUND_HALF_UP)


def _clp(value: Any) -> str:
    if isinstance(value, int):
        return "$" + f"{value:,}".replace(",", ".")
    return str(value)


def _rut(rut: Any) -> str:
    from origenlab_api.v2.tenders.bidder import format_rut

    try:
        return format_rut(str(rut))
    except ValueError:
        return str(rut or "")


def offer_docx(code: str, payload: dict[str, Any], bidder: dict[str, Any]) -> bytes:
    """«Oferta económica <code>» when the bases bring no offer form of their own: OrigenLab, the
    bidder's name and RUT, the quote's number and date, and the table of lines with neto, IVA and
    total (spec §6.2)."""
    from docx import Document

    doc = Document()
    doc.add_heading(f"Oferta económica {code}", level=1)
    doc.add_paragraph("OrigenLab")
    doc.add_paragraph(f"Oferente: {bidder.get('legal_name') or ''} · RUT {_rut(bidder.get('rut'))}")
    fecha = payload.get("fecha")
    shown = "-".join(reversed(fecha.split("-"))) if isinstance(fecha, str) else ""
    doc.add_paragraph(f"Cotización N° {payload.get('numero') or ''} · Fecha {shown}")
    header = ("Línea", "Descripción", "Cantidad", "Precio unitario neto", "Neto", "IVA", "Total")
    table = doc.add_table(rows=1, cols=len(header))
    for cell, text in zip(table.rows[0].cells, header, strict=True):
        cell.text = text
    for n, line in enumerate(payload.get("lineas") or [], start=1):
        values = (str(n), line["descripcion"], str(line["cantidad"]), _clp(line["precio_unitario_neto"]),
                  _clp(line["neto"]), _clp(line["iva"]), _clp(line["total"]))
        for cell, text in zip(table.add_row().cells, values, strict=True):
            cell.text = text
    totals = doc.add_table(rows=0, cols=2)
    for label, key in (("Total neto", "neto"), ("IVA", "iva"), ("Total", "total")):
        row = totals.add_row().cells
        row[0].text, row[1].text = label, _clp(payload.get(key))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ── the package ─────────────────────────────────────────────────────────────────────────────


def _safe(text: str) -> str:
    text = _CONTROL.sub("", text or "")
    text = re.split(r"[/\\]", text)[-1]
    text = _UNSAFE.sub("_", text)
    return " ".join(text.split()).strip(" .")


def safe_filename(filename: str, fallback: str) -> str:
    """A file name safe in a zip and on Windows, at most `MAX_NAME` characters, its extension kept."""
    name = _safe(filename) or _safe(fallback) or "documento"
    stem, dot, ext = name.rpartition(".")
    if not dot or not stem or len(ext) > 10:
        stem, ext = name, ""
    room = MAX_NAME - (len(ext) + 1 if ext else 0)
    return stem[:room].rstrip(" .") + (f".{ext}" if ext else "")


def package_name(code: str, filename: str, index: int) -> str:
    """`NN_<filename>`: the bases' own file name, sanitised, numbered in the checklist's order."""
    return f"{index:02d}_" + safe_filename(filename, f"documento_{code}")


def package_filename(code: str) -> str:
    return safe_filename(f"Oferta_{code}.zip", "Oferta.zip")


def write_package(target: BinaryIO, members: Iterable[tuple[str, Callable[[], bytes]]], code: str) -> None:
    """Write the zip into `target`, reading each member only when it is written, so one member at a
    time is in memory. Raises :class:`PackageTooLarge` past `MAX_PACKAGE_BYTES`."""
    written = 0
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.comment = f"Oferta {code}".encode()[:200]
        for index, (filename, read) in enumerate(members, start=1):
            data = read()
            written += len(data)
            if written > MAX_PACKAGE_BYTES:
                raise PackageTooLarge(str(written))
            z.writestr(package_name(code, filename, index), data)


def build_package(files: list[tuple[str, bytes]], code: str) -> bytes:
    """The zip of `files` (filename, bytes) in order, members named `NN_<filename>`; the output is
    named `package_filename(code)`. Raises :class:`PackageTooLarge` before writing anything."""
    if sum(len(data) for _, data in files) > MAX_PACKAGE_BYTES:
        raise PackageTooLarge("files")
    out = io.BytesIO()
    write_package(out, [(name, (lambda d=data: d)) for name, data in files], code)
    return out.getvalue()


__all__ = ["CHECKLIST_STATES", "CONTENT_TYPES", "EXTENSIONS", "MAX_ATTACHMENT_BYTES", "MAX_PACKAGE_BYTES",
           "OUTPUT_KINDS", "PackageTooLarge", "bidder_kinds_for", "build_package", "checklist_status",
           "is_datasheet_entry", "is_offer_entry", "offer_docx", "offer_payload", "package_filename", "package_name",
           "safe_filename", "write_package"]
