"""Find the fields of a buyer's form (docx, xlsx, fillable PDF), write values into them, or build a
generated .docx when the original is flat (spec 2026-10-19 §6, `fill` phase).

The form is an untrusted third-party file: its size is bounded before it is opened, the number of
slots is capped, and nothing here raises on a malformed file — `inspect_form` answers
`not_fillable` with a reason the caller can surface. A legacy .doc, a flat PDF or any file with no
blank to fill is `not_fillable` (ruling R13: the caller turns that into the checklist line
«Falta: convertir a Word (.docx)»); no function returns empty bytes silently.
"""

from __future__ import annotations

import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any, Final

MAX_FORM_BYTES: int = 20 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES: int = 200 * 1024 * 1024
MAX_SLOTS: int = 400
MAX_LABEL: Final = 150
MAX_CONTEXT: Final = 240
_MAX_SHEET_ROWS: Final = 2000
_MAX_SHEET_COLS: Final = 60

COMPLETAR_RE: Final = re.compile(r"\[completar: [^\]]+\]")
GENERATED_NOTE: Final = "El original no tiene campos rellenables; este documento reproduce su contenido llenado."
_BLANK = re.compile(r"_{3,}|\.{5,}|…{3,}")
_CONTROL = re.compile(r"[\x00-\x08\x0b-\x1f\x7f]")

NOT_FILLABLE: Final = "not_fillable"


class FormError(Exception):
    """A form could not be written; the message is for logs, never shown from model text."""


@dataclass(frozen=True)
class Slot:
    slot_id: str
    label: str
    context: str


@dataclass
class FormInspection:
    kind: str
    slots: list[Slot] = field(default_factory=list)
    reason: str = ""


def _clip(text: str, limit: int) -> str:
    return " ".join(_CONTROL.sub(" ", text).split())[:limit]


def _label_of(text: str) -> str:
    return _clip(text, MAX_LABEL).rstrip(" :.-–").strip()


def _not_fillable(reason: str) -> FormInspection:
    return FormInspection(NOT_FILLABLE, [], reason)


def _zip_ok(data: bytes) -> str:
    """"" when `data` is a zip whose members stay within bounds, else the reason."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
            if len(infos) > 5000 or sum(i.file_size for i in infos) > MAX_UNCOMPRESSED_BYTES:
                return "archivo_demasiado_grande"
    except (zipfile.BadZipFile, OSError, ValueError):
        return "archivo_ilegible"
    return ""


# ── docx ────────────────────────────────────────────────────────────────────────────────────


def _docx_cells(table: Any) -> list[list[Any]]:
    return [list(row.cells) for row in table.rows]


def _docx_slots(data: bytes) -> list[Slot]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    slots: list[Slot] = []
    for pi, para in enumerate(doc.paragraphs):
        matches = list(_BLANK.finditer(para.text))
        for n, m in enumerate(matches):
            before = para.text[: m.start()]
            if n:
                before = before[matches[n - 1].end():]
            label = _label_of(before) or _label_of(doc.paragraphs[pi - 1].text if pi else "") or "campo"
            slots.append(Slot(f"p{pi}" if n == 0 else f"p{pi}_{n + 1}", label, _clip(para.text, MAX_CONTEXT)))
    for ti, table in enumerate(doc.tables):
        for ri, cells in enumerate(_docx_cells(table)):
            row_text = _clip(" | ".join(c.text for c in cells), MAX_CONTEXT)
            seen: set[int] = set()
            for ci, cell in enumerate(cells):
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                text = cell.text.strip()
                if ci == 0 or (text and not _BLANK.fullmatch(text)):
                    continue
                left = cells[ci - 1]
                if left._tc is cell._tc or not left.text.strip() or _BLANK.fullmatch(left.text.strip()):
                    continue
                slots.append(Slot(f"t{ti}r{ri}c{ci}", _label_of(left.text) or "campo", row_text))
    return slots


def _docx_rewrite(para: Any, segments: list[tuple[str, bool]]) -> None:
    from docx.enum.text import WD_COLOR_INDEX

    runs = list(para.runs)
    for extra in runs[1:]:
        extra._element.getparent().remove(extra._element)
    first = runs[0] if runs else para.add_run("")
    first.text = ""
    for i, (text, highlight) in enumerate(segments):
        run = first if i == 0 else para.add_run()
        run.text = text
        if highlight:
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _docx_apply(data: bytes, values: dict[str, str]) -> bytes:
    from docx import Document

    doc = Document(io.BytesIO(data))
    for pi, para in enumerate(doc.paragraphs):
        text = para.text
        matches = list(_BLANK.finditer(text))
        if not matches:
            continue
        segments: list[tuple[str, bool]] = []
        cursor, changed = 0, False
        for n, m in enumerate(matches):
            value = values.get(f"p{pi}" if n == 0 else f"p{pi}_{n + 1}")
            if value is None:
                continue
            segments.append((text[cursor: m.start()], False))
            segments.append((value, True))
            cursor, changed = m.end(), True
        if changed:
            segments.append((text[cursor:], False))
            _docx_rewrite(para, [s for s in segments if s[0] or s[1]])
    for ti, table in enumerate(doc.tables):
        for ri, cells in enumerate(_docx_cells(table)):
            for ci, cell in enumerate(cells):
                value = values.get(f"t{ti}r{ri}c{ci}")
                if value is None or ci == 0:
                    continue
                para = cell.paragraphs[0]
                for other in cell.paragraphs[1:]:
                    other._element.getparent().remove(other._element)
                _docx_rewrite(para, [(value, True)])
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ── xlsx ────────────────────────────────────────────────────────────────────────────────────


def _xlsx_slots(data: bytes) -> list[Slot]:
    import openpyxl
    from openpyxl.utils import get_column_letter

    wb = openpyxl.load_workbook(io.BytesIO(data))
    slots: list[Slot] = []
    for ws in wb.worksheets:
        max_row, max_col = min(ws.max_row or 0, _MAX_SHEET_ROWS), min(ws.max_column or 0, _MAX_SHEET_COLS)
        merged = {(r, c): rng for rng in ws.merged_cells.ranges
                  for r in range(rng.min_row, min(rng.max_row, max_row) + 1)
                  for c in range(rng.min_col, min(rng.max_col, max_col) + 1)}

        def free(r: int, c: int) -> bool:
            rng = merged.get((r, c))
            if rng is not None and (r, c) != (rng.min_row, rng.min_col):
                return False
            return ws.cell(r, c).value in (None, "")

        claimed: set[tuple[int, int]] = set()
        for r in range(1, max_row + 1):
            row_text = _clip(" | ".join(str(ws.cell(r, c).value) for c in range(1, max_col + 1)
                                        if ws.cell(r, c).value not in (None, "")), MAX_CONTEXT)
            for c in range(1, max_col + 1):
                value = ws.cell(r, c).value
                if not isinstance(value, str) or not value.strip() or len(value) > 120 or not re.search(r"[^\W\d_]", value):
                    continue
                rng = merged.get((r, c))
                right = ((rng.max_col if rng else c) + 1)
                below = ((rng.max_row if rng else r) + 1)
                neighbours = [ws.cell(r, c - 1).value if c > 1 else None, ws.cell(r, right).value]
                header = any(isinstance(v, str) and v.strip() for v in neighbours)
                targets = ((below, c), (r, right)) if header else ((r, right), (below, c))
                for tr, tc in targets:
                    if tr <= _MAX_SHEET_ROWS and tc <= _MAX_SHEET_COLS and free(tr, tc) and (tr, tc) not in claimed:
                        claimed.add((tr, tc))
                        slots.append(Slot(f"{ws.title}!{get_column_letter(tc)}{tr}", _label_of(value) or "campo", row_text))
                        break
    return slots


def _xlsx_apply(data: bytes, values: dict[str, str]) -> bytes:
    import openpyxl
    from openpyxl.styles import PatternFill

    wb = openpyxl.load_workbook(io.BytesIO(data))
    for slot_id, value in values.items():
        sheet, _, ref = slot_id.rpartition("!")
        if sheet not in wb.sheetnames:
            continue
        cell = wb[sheet][ref]
        cell.value = value
        cell.fill = PatternFill("solid", fgColor="FFFF00")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ── fillable PDF ────────────────────────────────────────────────────────────────────────────


def _pdf_slots(data: bytes) -> list[Slot]:
    import fitz

    slots: list[Slot] = []
    seen: set[str] = set()
    with fitz.open(stream=data, filetype="pdf") as doc:
        if doc.needs_pass or not doc.is_form_pdf:
            return []
        for page in doc:
            for w in page.widgets() or []:
                name = w.field_name
                if not name or name in seen or w.field_type_string not in ("Text", "ComboBox"):
                    continue
                seen.add(name)
                r = w.rect
                near = page.get_textbox(fitz.Rect(max(r.x0 - 260, 0), r.y0 - 3, r.x0, r.y1 + 3))
                label = _label_of(w.field_label or near) or _label_of(name)
                slots.append(Slot(name, label, _clip(near, MAX_CONTEXT)))
    return slots


def _pdf_apply(data: bytes, values: dict[str, str]) -> bytes:
    import fitz

    with fitz.open(stream=data, filetype="pdf") as doc:
        for page in doc:
            for w in page.widgets() or []:
                if w.field_name in values:
                    w.field_value = values[w.field_name]
                    w.update()
        return doc.tobytes()


# ── public API ──────────────────────────────────────────────────────────────────────────────


def inspect_form(filename: str, data: bytes) -> FormInspection:
    """Kind, slots and (when not fillable) the reason. Never raises."""
    name = (filename or "").lower()
    if not data:
        return _not_fillable("archivo_vacio")
    if len(data) > MAX_FORM_BYTES:
        return _not_fillable("archivo_demasiado_grande")
    try:
        if name.endswith((".docx", ".xlsx")):
            problem = _zip_ok(data)
            if problem:
                return _not_fillable(problem)
            kind = "docx" if name.endswith(".docx") else "xlsx"
            slots = _docx_slots(data) if kind == "docx" else _xlsx_slots(data)
        elif name.endswith(".pdf"):
            if not data.lstrip()[:5].startswith(b"%PDF"):
                return _not_fillable("archivo_ilegible")
            kind, slots = "pdf_fillable", _pdf_slots(data)
        elif name.endswith((".doc", ".xls")):
            return _not_fillable("formato_antiguo")
        else:
            return _not_fillable("formato_no_soportado")
    except Exception:  # noqa: BLE001 - untrusted file: any parser failure is "not fillable"
        return _not_fillable("archivo_ilegible")
    if not slots:
        return _not_fillable("sin_campos")
    return FormInspection(kind, slots[:MAX_SLOTS])


def form_slots(filename: str, data: bytes) -> tuple[str, list[Slot]]:
    result = inspect_form(filename, data)
    return result.kind, result.slots


def apply_values(kind: str, data: bytes, values: dict[str, str]) -> bytes:
    """The form with `values` written in (docx/xlsx runs highlighted yellow). Raises
    :class:`FormError` rather than return empty or unchanged-by-failure bytes."""
    writers = {"docx": _docx_apply, "xlsx": _xlsx_apply, "pdf_fillable": _pdf_apply}
    writer = writers.get(kind)
    if writer is None:
        raise FormError(f"cannot fill a {kind} form")
    try:
        out = writer(data, values)
    except Exception as exc:  # noqa: BLE001
        raise FormError(f"form could not be written: {type(exc).__name__}") from None
    if not out:
        raise FormError("empty output")
    return out


def generated_docx(title: str, sections: list[tuple[str, str]], source_filename: str) -> bytes:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    doc = Document()
    doc.add_paragraph(GENERATED_NOTE)
    doc.add_paragraph(f"Archivo original: {_clip(source_filename, 200)}")
    doc.add_heading(_clip(title, 200) or "Documento", level=1)
    for heading, body in sections:
        doc.add_heading(_clip(heading, 200), level=2)
        for line in (body or "").split("\n"):
            para = doc.add_paragraph()
            cursor = 0
            for m in COMPLETAR_RE.finditer(line):
                para.add_run(line[cursor: m.start()])
                para.add_run(m.group(0)).font.highlight_color = WD_COLOR_INDEX.YELLOW
                cursor = m.end()
            para.add_run(line[cursor:])
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def count_completar(data: bytes, kind: str) -> int:
    """How many «[completar: …]» marks a written form carries (0 for an unreadable file)."""
    try:
        if kind == "docx":
            from docx import Document

            doc = Document(io.BytesIO(data))
            texts = [p.text for p in doc.paragraphs]
            texts += [c.text for t in doc.tables for row in t.rows for c in row.cells]
        elif kind == "xlsx":
            import openpyxl

            wb = openpyxl.load_workbook(io.BytesIO(data))
            texts = [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
        elif kind == "pdf_fillable":
            import fitz

            with fitz.open(stream=data, filetype="pdf") as doc:
                texts = [str(w.field_value or "") for page in doc for w in page.widgets() or []]
        else:
            return 0
    except Exception:  # noqa: BLE001
        return 0
    return sum(len(COMPLETAR_RE.findall(t)) for t in texts)


__all__ = ["COMPLETAR_RE", "FormError", "FormInspection", "GENERATED_NOTE", "Slot", "apply_values", "count_completar",
           "form_slots", "generated_docx", "inspect_form"]
