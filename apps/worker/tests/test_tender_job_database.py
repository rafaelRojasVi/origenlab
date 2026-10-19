"""The tender `read` job on a disposable `origenlab_test_<hex>` database, as the real
`origenlab_worker` login, with the anexos in an in-process S3 (moto): a queued preparation is
claimed, read and written back with its spend; the monthly cap stops it before the model is
called; a refusal fails it by code; a locked row is skipped.

The tests share one database and one S3, so each seeds its own case. Every value is invented."""

from __future__ import annotations

import hashlib
import io
import json
import re
import uuid
from datetime import UTC, datetime
from types import SimpleNamespace

import psycopg
import pytest

from conftest import REGION
from dbhelp import owner_rows
from origenlab_worker.database import local_test_target, open_verified_connection
from origenlab_worker.storage import S3EmlStore
from origenlab_worker.tender.job import (
    BID_REQUIRED_POLICIES,
    BID_REQUIRED_PRIVILEGES,
    DOCUMENTS_BUCKET,
    BidModels,
    BidPreparer,
    ConnectionLost,
    run_bid_sweep,
)
from origenlab_worker.tender.read_model import RESUMEN_FIELDS, CallInterrupted
from test_quote_reader import _minimal_pdf
from v2_command_harness import build_disposable_database, needs_worker_db, worker_dsn

pytestmark = needs_worker_db
MODELS = BidModels(read="claude-opus-5-5", fill="claude-opus-5-5", match="claude-sonnet-5-5")
#: The cap test's own month: its seeded spend never reaches the other tests' month.
CAP_MONTH = datetime(2026, 2, 10, 15, 0, tzinfo=UTC)


def _now() -> datetime:
    return datetime.now(UTC)


@pytest.fixture(scope="module")
def dsn():
    yield from build_disposable_database()


@pytest.fixture(scope="module")
def store(s3_client):
    s3_client.create_bucket(Bucket=DOCUMENTS_BUCKET, CreateBucketConfiguration={"LocationConstraint": REGION})
    return S3EmlStore(s3_client, bucket=DOCUMENTS_BUCKET)


@pytest.fixture(scope="module")
def operator(dsn):
    return owner_rows(dsn, "insert into platform.operator (auth_user_id, email_norm, display_name, role, status) "
                           "values (gen_random_uuid(), 'op-licitaciones@example.test', 'Op', 'sales', 'active') "
                           "returning id::text")[0][0]


def _conn(dsn):
    return open_verified_connection(local_test_target(worker_dsn(dsn)), application_name="test-bid-preparation",
                                    privileges=BID_REQUIRED_PRIVILEGES, policies=BID_REQUIRED_POLICIES)


def _case(dsn, operator) -> str:
    return owner_rows(dsn, "insert into crm.opportunity (title, stage, owner_operator_id, kind, tender_status) "
                           "values (%s, 'lead', %s, 'tender', 'preparando') returning id::text",
                      (f"Licitación ficticia {uuid.uuid4().hex[:6]}", operator))[0][0]


def _annex(dsn, store, operator, case: str, data: bytes, *, filename: str = "bases.pdf",
           content_type: str = "application/pdf", ext: str = "pdf") -> str:
    sha = hashlib.sha256(data).hexdigest()
    path = f"opportunities/{case}/{sha}.{ext}"
    store._client.put_object(Bucket=DOCUMENTS_BUCKET, Key=path, Body=data)
    return owner_rows(dsn, "insert into crm.opportunity_document (opportunity_id, kind, storage_path, sha256, "
                           "content_type, byte_size, original_filename, source, attached_by_operator_id) "
                           "values (%s, 'tender_annex', %s, %s, %s, %s, %s, 'upload', %s) returning id::text",
                      (case, path, sha, content_type, len(data), filename, operator))[0][0]


def _prep(dsn, operator, case: str, annexes: list[str]) -> str:
    return owner_rows(dsn, "insert into evidence.bid_preparation (opportunity_id, phase, annex_document_ids, "
                           "requested_by_operator_id) values (%s, 'read', %s::uuid[], %s) returning id::text",
                      (case, annexes, operator))[0][0]


def _queued(dsn, store, operator, text: str = "Fecha de cierre 30 de octubre") -> tuple[str, str, str]:
    case = _case(dsn, operator)
    doc = _annex(dsn, store, operator, case, _minimal_pdf(f"{text} {uuid.uuid4().hex[:6]}"))
    return case, doc, _prep(dsn, operator, case, [doc])


def _answer(doc: str) -> dict:
    cita = {"archivo_id": doc, "ubicacion": "p.1"}
    return {
        "resumen": {name: {"texto": f"Texto de {name}", "citas": [cita]} for name in RESUMEN_FIELDS},
        "checklist": [{"documento": "Anexo 1", "tipo": "formulario", "requiere_firma": True,
                       "archivo_formulario_id": doc, "fuente": cita}],
        "items": [{"linea": 1, "descripcion": "Balanza analítica", "cantidad": "2", "unidad": "unidad",
                   "requisitos": ["0,1 mg"], "citas": [cita]}],
    }


class Reader:
    def __init__(self, answer=None, *, stop="end_turn", error=None, usage=None, model=None, stop_details=None,
                 during=None):
        self.answer, self.stop, self.error, self.usage = answer, stop, error, usage
        self.model, self.stop_details, self.during = model, stop_details, during
        self.calls = 0
        self.requests: list[dict] = []

    def __call__(self, request):
        self.calls += 1
        self.requests.append(request)
        if self.during is not None:
            self.during()
        if self.error is not None:
            raise self.error
        return SimpleNamespace(stop_reason=self.stop, model=self.model or request["model"],
                               stop_details=self.stop_details,
                               content=[SimpleNamespace(type="text", text=json.dumps(self.answer or {}))],
                               usage=self.usage or SimpleNamespace(input_tokens=1000, output_tokens=200,
                                                                   cache_read_input_tokens=0,
                                                                   cache_creation_input_tokens=0))


def _status_error(status: int):
    import anthropic
    import httpx

    response = httpx.Response(status, request=httpx.Request("POST", "https://api.example.test/v1/messages"))
    return anthropic.APIStatusError("refused", response=response, body=None)


def _usage_rows(dsn, prep: str):
    return owner_rows(dsn, "select model, input_tokens, output_tokens, usd_cost::text from platform.model_usage "
                           "where bid_preparation_id = %s order by created_at, id", (prep,))


def _row(dsn, prep: str):
    return owner_rows(dsn, "select status, failure_code, result, models, started_at is not null, "
                           "finished_at is not null from evidence.bid_preparation where id = %s", (prep,))[0]


def _usage(dsn, prep: str, step: str = "read") -> int:
    return owner_rows(dsn, "select count(*) from platform.model_usage where bid_preparation_id = %s and step = %s",
                      (prep, step))[0][0]


def test_job_ready_and_usage_recorded(dsn, store, operator) -> None:
    case, doc, prep = _queued(dsn, store, operator)
    reader = Reader(_answer(doc))
    with _conn(dsn) as conn:
        status = BidPreparer(conn, store, reader, MODELS, now=_now).run(uuid.UUID(prep))
    assert status == "ready" and _usage(dsn, prep) == 1 and reader.calls == 1
    status, code, result, models, started, finished = _row(dsn, prep)
    assert (status, code, started, finished) == ("ready", None, True, True)
    assert set(result) == {"resumen", "checklist", "items", "annexes", "match"}
    assert result["annexes"] == [{"document_id": doc, "filename": "bases.pdf", "outcome": "extraction_success"}]
    assert result["resumen"]["cierre"]["sin_cita"] is False
    assert result["items"][0]["descripcion"] == "Balanza analítica"
    assert models == {"read": "claude-opus-5-5"}
    content = reader.requests[0]["messages"][0]["content"]
    assert f"<<<ARCHIVO id={doc} nombre=bases.pdf>>>" in content and "Fecha de cierre" in content


def test_job_stops_at_cap_before_calling(dsn, store, operator) -> None:
    case, doc, prep = _queued(dsn, store, operator)
    owner_rows(dsn, "insert into platform.model_usage (feature, model, step, usd_cost, created_at) "
                    "values ('licitaciones', 'claude-opus-5-5', 'read', 59.99, %s)", (CAP_MONTH,))
    reader = Reader(_answer(doc))
    with _conn(dsn) as conn:
        assert BidPreparer(conn, store, reader, MODELS, now=lambda: CAP_MONTH).run(uuid.UUID(prep)) == "budget_stopped"
    assert reader.calls == 0 and _usage(dsn, prep) == 0
    assert _row(dsn, prep)[:2] == ("budget_stopped", "budget_reached")


def test_refusal_stop_reason_fails_row(dsn, store, operator) -> None:
    case, doc, prep = _queued(dsn, store, operator)
    with _conn(dsn) as conn:
        assert BidPreparer(conn, store, Reader(stop="refusal"), MODELS, now=_now).run(uuid.UUID(prep)) == "failed"
    assert _row(dsn, prep)[:2] == ("failed", "model_refused")
    assert _usage(dsn, prep) == 1  # a refused call is still recorded: it may be billed


def test_a_truncated_answer_fails_the_row(dsn, store, operator) -> None:
    case, doc, prep = _queued(dsn, store, operator)
    with _conn(dsn) as conn:
        BidPreparer(conn, store, Reader(stop="max_tokens"), MODELS, now=_now).run(uuid.UUID(prep))
    assert _row(dsn, prep)[:2] == ("failed", "model_truncated")


def test_an_api_status_error_fails_the_row_without_usage(dsn, store, operator) -> None:
    case, doc, prep = _queued(dsn, store, operator)
    with _conn(dsn) as conn:
        BidPreparer(conn, store, Reader(error=_status_error(529)), MODELS, now=_now).run(uuid.UUID(prep))
    assert _row(dsn, prep)[:2] == ("failed", "request_failed") and _usage(dsn, prep) == 0


@pytest.mark.parametrize("error", [CallInterrupted(), TimeoutError("read timed out")])
def test_a_call_lost_after_it_was_sent_records_every_attempt_at_its_worst_case(dsn, store, operator, error) -> None:
    """With fallbacks on, the requested model may decline mid-stream and the fallback start before
    the break: both are billed, so both are recorded, and the total is what the cap check assumed."""
    from decimal import Decimal

    from origenlab_worker.tender.job import call_worst_case

    case, doc, prep = _queued(dsn, store, operator)
    reader = Reader(error=error)
    with _conn(dsn) as conn:
        BidPreparer(conn, store, reader, MODELS, now=_now).run(uuid.UUID(prep))
    assert _row(dsn, prep)[:2] == ("failed", "request_failed")
    rows = _usage_rows(dsn, prep)
    assert len(rows) == 2 and rows[0][0] != rows[1][0] and "claude-opus-5-5" in {r[0] for r in rows}
    assert all(r[2] == 32000 for r in rows)
    assert sum(Decimal(r[3]) for r in rows) == call_worst_case(reader.requests[0], MODELS.read)


def test_every_billed_attempt_of_a_fallback_is_recorded(dsn, store, operator) -> None:
    case, doc, prep = _queued(dsn, store, operator)
    usage = SimpleNamespace(input_tokens=1000, output_tokens=300, iterations=[
        SimpleNamespace(type="message", model=None, input_tokens=1000, output_tokens=100,
                        cache_read_input_tokens=0, cache_creation_input_tokens=0),
        SimpleNamespace(type="fallback_message", model="claude-opus-4-8", input_tokens=1000, output_tokens=300,
                        cache_read_input_tokens=0, cache_creation_input_tokens=0)])
    reader = Reader(_answer(doc), usage=usage, model="claude-opus-4-8")
    with _conn(dsn) as conn:
        assert BidPreparer(conn, store, reader, MODELS, now=_now).run(uuid.UUID(prep)) == "ready"
    assert sorted(r[:3] for r in _usage_rows(dsn, prep)) == [("claude-opus-4-8", 1000, 300),
                                                             ("claude-opus-5-5", 1000, 100)]
    assert _row(dsn, prep)[3] == {"read": "claude-opus-4-8"}


def test_the_refusal_category_is_kept_on_the_failed_row(dsn, store, operator) -> None:
    case, doc, prep = _queued(dsn, store, operator)
    reader = Reader(stop="refusal", stop_details=SimpleNamespace(category="bio"))
    with _conn(dsn) as conn:
        BidPreparer(conn, store, reader, MODELS, now=_now).run(uuid.UUID(prep))
    assert _row(dsn, prep)[:3] == ("failed", "model_refused", {"refusal_category": "bio"})


def test_a_connection_lost_during_the_call_records_usage_on_a_fresh_one_and_stops(dsn, store, operator) -> None:
    case, doc, prep = _queued(dsn, store, operator)
    with _conn(dsn) as conn:
        reader = Reader(_answer(doc), during=conn.close)  # the session dies while the model answers
        preparer = BidPreparer(conn, store, reader, MODELS, now=_now, reconnect=lambda: _conn(dsn))
        with pytest.raises(ConnectionLost):
            preparer.run(uuid.UUID(prep))
    assert _usage(dsn, prep) == 1
    assert _row(dsn, prep)[:2] == ("failed", "connection_lost")


def test_the_sweep_stops_when_the_connection_is_lost(dsn, store, operator) -> None:
    case, doc, prep = _queued(dsn, store, operator)
    with _conn(dsn) as conn:
        reader = Reader(_answer(doc), during=conn.close)
        counts = run_bid_sweep(conn, store, reader, MODELS, now=_now, reconnect=lambda: _conn(dsn))
    assert counts["outcome"] == "connection_lost" and counts["claimed"] == 1
    assert _row(dsn, prep)[:2] == ("failed", "connection_lost")


def test_an_unexpected_queue_error_is_reported_by_class_only(dsn, store, operator, monkeypatch) -> None:
    def boom(self):
        raise psycopg.DataError("valor secreto de una fila")

    monkeypatch.setattr(BidPreparer, "claim", boom)
    with _conn(dsn) as conn:
        counts = run_bid_sweep(conn, store, Reader(), MODELS, now=_now)
    assert counts["outcome"] == "error" and counts["error"] == "DataError"
    assert "secreto" not in json.dumps(counts)


def test_no_readable_annex_fails_before_the_model(dsn, store, operator) -> None:
    case = _case(dsn, operator)
    doc = _annex(dsn, store, operator, case, b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600,
                 filename="anexo.doc", content_type="application/msword", ext="doc")
    prep = _prep(dsn, operator, case, [doc])
    reader = Reader(_answer(doc))
    with _conn(dsn) as conn:
        assert BidPreparer(conn, store, reader, MODELS, now=_now).run(uuid.UUID(prep)) == "failed"
    assert _row(dsn, prep)[:2] == ("failed", "no_readable_annex") and reader.calls == 0


def test_a_legacy_word_annex_beside_a_readable_one_is_a_checklist_item(dsn, store, operator) -> None:
    case = _case(dsn, operator)
    pdf = _annex(dsn, store, operator, case, _minimal_pdf(f"Bases {uuid.uuid4().hex[:6]}"))
    doc = _annex(dsn, store, operator, case, b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x01" * 600,
                 filename="anexo 3.doc", content_type="application/msword", ext="doc")
    prep = _prep(dsn, operator, case, [pdf, doc])
    with _conn(dsn) as conn:
        assert BidPreparer(conn, store, Reader(_answer(pdf)), MODELS, now=_now).run(uuid.UUID(prep)) == "ready"
    result = _row(dsn, prep)[2]
    assert [a["outcome"] for a in result["annexes"]] == ["extraction_success", "needs_converter"]
    [item] = [c for c in result["checklist"] if c.get("archivo_formulario_id") == doc]
    assert item["documento"] == "Falta: convertir a Word (.docx): anexo 3.doc"


def test_claim_skips_locked(dsn, store, operator) -> None:
    _, _, first = _queued(dsn, store, operator)
    _, _, second = _queued(dsn, store, operator)
    with psycopg.connect(dsn) as holder, holder.cursor() as cur:
        cur.execute("set role origenlab_owner")
        cur.execute("select id from evidence.bid_preparation where id = %s for update", (first,))
        with _conn(dsn) as conn:
            claimed = BidPreparer(conn, store, Reader(), MODELS, now=_now).claim()
        holder.rollback()
    assert str(claimed) == second
    assert _row(dsn, second)[0] == "running" and _row(dsn, first)[0] == "queued"
    with _conn(dsn) as conn:  # leave nothing queued or running for the other tests
        preparer = BidPreparer(conn, store, Reader(error=RuntimeError("x")), MODELS, now=_now)
        assert preparer.claim() == uuid.UUID(first)
        assert preparer.claim() is None
        preparer.run(uuid.UUID(first))
        preparer.run(uuid.UUID(second))


def test_the_sweep_takes_the_lock_recovers_an_interrupted_row_and_runs_the_queue(dsn, store, operator) -> None:
    case, doc, stuck = _queued(dsn, store, operator)
    owner_rows(dsn, "update evidence.bid_preparation set status = 'running', started_at = now() where id = %s",
               (stuck,))
    case2, doc2, prep = _queued(dsn, store, operator)
    with _conn(dsn) as conn:
        counts = run_bid_sweep(conn, store, Reader(_answer(doc2)), MODELS, now=_now)
    assert counts == {"outcome": "ran", "interrupted": 1, "claimed": 1, "statuses": {"ready": 1}}
    assert _row(dsn, stuck)[:2] == ("failed", "interrupted") and _row(dsn, prep)[0] == "ready"


def test_a_second_sweep_while_one_holds_the_lock_does_nothing(dsn, store, operator) -> None:
    from origenlab_worker.tender.job import LOCK_NAME

    with _conn(dsn) as holder:
        holder.execute("select pg_advisory_lock(hashtextextended(%s, 0))", (LOCK_NAME,))
        with _conn(dsn) as conn:
            assert run_bid_sweep(conn, store, Reader(), MODELS, now=_now) == {"outcome": "locked"}


# ── 12a: the match step at the end of `read` ────────────────────────────────────────────────


def _maker(dsn) -> str:
    return owner_rows(dsn, "insert into crm.organization (kind, name, confirmation) values ('manufacturer', %s, "
                           "'machine_proposed') returning id::text", (f"Fabricante Ficticio {uuid.uuid4().hex[:6]}",))[0][0]


def _product(dsn, model: str, name: str = "Balanza analítica") -> str:
    return owner_rows(dsn, "insert into catalog.product (manufacturer_organization_id, model_number, name) "
                           "values (%s, %s, %s) returning id::text", (_maker(dsn), model, name))[0][0]


class Script:
    """A reader answering by the request: the match step, a form with fields, a flat form."""

    def __init__(self, read_answer=None, *, match=None, value="Oferente Ejemplo", sections=None):
        self.read_answer, self.match, self.value = read_answer, match, value
        self.sections = sections if sections is not None else [
            {"titulo": "Declaración", "texto": "Yo, Oferente Ejemplo, declaro lo que indican las bases."}]
        self.calls = 0
        self.requests: list[dict] = []

    def answer(self, request):
        content = request["messages"][0]["content"]
        if content.startswith("ITEMS") or "<<<ITEMS>>>" in content:
            return self.match(request) if callable(self.match) else self.match
        if "OFERENTE" in content and "CAMPOS" in content:
            slots = re.findall(r"^\[\[(.+?)\]\] etiqueta:", content, re.M)
            return {"valores": [{"slot_id": s, "valor": self.value} for s in slots],
                    "secciones": [] if slots else self.sections}
        return self.read_answer

    def __call__(self, request):
        self.calls += 1
        self.requests.append(request)
        return SimpleNamespace(stop_reason="end_turn", model=request["model"], stop_details=None,
                               content=[SimpleNamespace(type="text", text=json.dumps(self.answer(request)))],
                               usage=SimpleNamespace(input_tokens=1000, output_tokens=200,
                                                     cache_read_input_tokens=0, cache_creation_input_tokens=0))


def _steps(dsn, prep: str) -> list[str]:
    return [r[0] for r in owner_rows(dsn, "select step from platform.model_usage where bid_preparation_id = %s "
                                          "order by created_at, id", (prep,))]


def test_the_read_matches_items_to_catalog_products_with_one_match_call(dsn, store, operator) -> None:
    model = f"PX{uuid.uuid4().int % 10**4:04d}Z"
    pid = _product(dsn, model)
    case, doc, prep = _queued(dsn, store, operator)
    answer = _answer(doc)
    answer["items"][0]["descripcion"] = f"Balanza analítica {model} 220 g"
    answer["items"].append({"linea": 2, "descripcion": "Gradilla", "cantidad": "1", "unidad": "", "requisitos": [],
                            "citas": []})
    script = Script(answer, match=lambda request: {"elecciones": [
        {"linea": 1, "product_id": pid, "motivo": "Mismo modelo"}]})
    with _conn(dsn) as conn:
        assert BidPreparer(conn, store, script, MODELS, now=_now).run(uuid.UUID(prep)) == "ready"
    result, models = _row(dsn, prep)[2], _row(dsn, prep)[3]
    first, second = result["items"]
    assert (first["product_id"], first["match_reason"]) == (pid, "Mismo modelo")
    assert [c["product_id"] for c in first["candidatos"]] == [pid]
    assert (second["product_id"], second["candidatos"]) == (None, [])
    assert result["match"] == {"estado": "hecho"}
    assert _steps(dsn, prep) == ["read", "match"] and models == {"read": MODELS.read, "match": MODELS.match}
    match_request = script.requests[1]
    assert match_request["model"] == "claude-sonnet-5-5" and match_request["output_config"]["effort"] == "low"


def test_no_catalog_candidate_means_no_match_call(dsn, store, operator) -> None:
    case, doc, prep = _queued(dsn, store, operator)
    answer = _answer(doc)
    answer["items"][0]["descripcion"] = "Gradilla plástica"
    script = Script(answer)
    with _conn(dsn) as conn:
        assert BidPreparer(conn, store, script, MODELS, now=_now).run(uuid.UUID(prep)) == "ready"
    result = _row(dsn, prep)[2]
    assert script.calls == 1 and _steps(dsn, prep) == ["read"]
    assert result["match"] == {"estado": "sin_candidatos"} and result["items"][0]["product_id"] is None


# ── the fill phase ──────────────────────────────────────────────────────────────────────────

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
CODE = "1234567-12-LE26"


def _docx(*paragraphs: str) -> bytes:
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _tender_case(dsn, operator) -> str:
    """A tender case with its Mercado Público notice (the code and the buyer)."""
    case = _case(dsn, operator)
    code = f"{uuid.uuid4().int % 10**7:07d}-12-LE26"
    notice = owner_rows(dsn, "insert into procurement.notice (codigo_externo, title, buyer_name, head) "
                             "values (%s, 'Compra de balanzas', 'Hospital Ejemplo', '{}'::jsonb) returning id::text", (code,))[0][0]
    owner_rows(dsn, "insert into crm.opportunity_evidence (opportunity_id, notice_id, relation, linked_by_operator_id) "
                    "values (%s, %s, 'origin', %s) returning id", (case, notice, operator))
    return case


def _offer(dsn, operator, case: str, lines: list[tuple[str, int, int, str | None]], *, status: str = "approved") -> str:
    quote = owner_rows(dsn, "insert into crm.quote (opportunity_id, quote_number, created_by_operator_id) "
                            "values (%s, %s, %s) returning id::text",
                       (case, f"{uuid.uuid4().int % 10**5:05d}-26", operator))[0][0]
    net = sum(q * p for _, q, p, _ in lines)
    revision = owner_rows(
        dsn, "insert into crm.quote_revision (quote_id, revision_no, status, quote_currency, price_decimals) "
             "values (%s, 1, 'draft', 'CLP', 0) returning id::text", (quote,))[0][0]
    for n, (description, quantity, price, product) in enumerate(lines, start=1):
        owner_rows(dsn, "insert into crm.quote_line (quote_revision_id, line_no, line_kind, description, quantity, "
                        "unit_price_qc, line_total_qc, product_id) values (%s, %s, 'item', %s, %s, %s, %s, %s) "
                        "returning id", (revision, n, description, quantity, price, quantity * price, product))
    if status in ("approved", "sent"):
        owner_rows(dsn, "update crm.quote_revision set status = 'approved', subtotal = %s, discount_total = 0, "
                        "tax_base = %s, tax_total = %s, grand_total = %s, totals_computed_at = now(), "
                        "approved_at = now(), approved_by_operator_id = %s, party_snapshot = '{}'::jsonb, "
                        "party_snapshot_version = 1 where id = %s returning id",
                   (net, net, round(net * 0.19), round(net * 1.19), operator, revision))
    return revision


def _bidder(dsn, operator) -> None:
    owner_rows(dsn, "insert into crm.bidder_profile (bidder_kind, legal_name, rut, comuna, email, set_by_operator_id, "
                    "reason) values ('natural_person', 'Oferente Ejemplo', '11111111-1', 'Santiago', "
                    "'oferente@example.test', %s, 'prueba') returning id", (operator,))


def _bidder_document(dsn, store, operator, kind: str, data: bytes, filename: str) -> str:
    sha = hashlib.sha256(data).hexdigest()
    path = f"bidder/{sha}.pdf"
    store._client.put_object(Bucket=DOCUMENTS_BUCKET, Key=path, Body=data)
    return owner_rows(dsn, "insert into crm.bidder_document (kind, title, storage_path, sha256, content_type, byte_size, "
                           "original_filename, added_by_operator_id) values (%s, %s, %s, %s, 'application/pdf', %s, %s, %s) "
                           "returning id::text", (kind, filename, path, sha, len(data), filename, operator))[0][0]


def _ready_read(dsn, operator, case: str, annexes: list[str], checklist: list[dict], items=None) -> str:
    result = {"resumen": {}, "checklist": checklist, "items": items or [], "annexes": []}
    return owner_rows(dsn, "insert into evidence.bid_preparation (opportunity_id, phase, status, annex_document_ids, "
                           "requested_by_operator_id, result, started_at, finished_at) values "
                           "(%s, 'read', 'ready', %s::uuid[], %s, %s, now(), now()) returning id::text",
                      (case, annexes, operator, psycopg.types.json.Jsonb(result)))[0][0]


def _fill_prep(dsn, operator, case: str, read: str, annexes: list[str], revision: str) -> str:
    return owner_rows(dsn, "insert into evidence.bid_preparation (opportunity_id, phase, annex_document_ids, "
                           "based_on_preparation_id, quote_revision_id, requested_by_operator_id) values "
                           "(%s, 'fill', %s::uuid[], %s, %s, %s) returning id::text",
                      (case, annexes, read, revision, operator))[0][0]


def _entry(documento: str, tipo: str, form: str | None = None, *, firma: bool = False) -> dict:
    return {"documento": documento, "tipo": tipo, "requiere_firma": firma, "archivo_formulario_id": form,
            "fuente": None, "sin_cita": True}


def _outputs(dsn, prep: str) -> list[dict]:
    return owner_rows(dsn, "select outputs from evidence.bid_preparation where id = %s", (prep,))[0][0]


class TenderDrive:
    """The archive's in-memory Drive with the generic upload the package needs."""

    def __new__(cls):
        from test_v2_quote_case_archive import FakeDrive

        class _Drive(FakeDrive):
            def upload_file(self, *, name, parent_id, app_properties, data, mime_type):
                fid = next(self.ids)
                self.writes.append(f"upload:{name}")
                return self.add(fid, name, parent_id, dict(app_properties), mime=mime_type, data=data)

        return _Drive()


def _drive_target():
    from test_v2_quote_case_archive import CASOS, ENV, PEND, PRINCIPAL

    from origenlab_worker.drive_filing import DriveTarget

    return DriveTarget(principal=PRINCIPAL, casos_folder_id=CASOS, protected_folder_ids=frozenset({PEND, ENV}))


def _case_folder(drive, case: str) -> str:
    from origenlab_api.v2 import quote_case_archive as qa
    from test_v2_quote_case_archive import CASOS

    folder = f"case-{uuid.uuid4().hex[:8]}"
    drive.add(folder, "Caso 01300 — Hospital Ejemplo", CASOS, {qa.PROP_KIND: qa.KIND_CASE, qa.PROP_CASE_KEY: case})
    return folder


def _full_fill(dsn, store, operator, drive):
    """A case with: bases (no form), a Word form with a blank, a flat PDF declaration, a Word economic
    offer form, a legacy .doc form, a line asking for the cédula, and one asking for spec sheets."""
    case = _tender_case(dsn, operator)
    bases = _annex(dsn, store, operator, case, _minimal_pdf(f"Bases {uuid.uuid4().hex[:6]}"))
    form = _annex(dsn, store, operator, case, _docx("Anexo 1", f"Nombre del oferente: ________ {uuid.uuid4().hex[:4]}"),
                  filename="Anexo 1 Identificación.docx", content_type=DOCX, ext="docx")
    flat = _annex(dsn, store, operator, case, _minimal_pdf(f"Declaracion jurada {uuid.uuid4().hex[:6]}"),
                  filename="Anexo 2 Declaración.pdf")
    offer_form = _annex(dsn, store, operator, case,
                        _docx("Formulario Oferta Económica", f"Precio total neto: ________ {uuid.uuid4().hex[:4]}"),
                        filename="Anexo 5 Oferta Económica.docx", content_type=DOCX, ext="docx")
    legacy = _annex(dsn, store, operator, case, b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + uuid.uuid4().bytes * 40,
                    filename="Anexo 3 Experiencia.doc", content_type="application/msword", ext="doc")
    annexes = [bases, form, flat, offer_form, legacy]
    checklist = [
        _entry("Anexo 1 Identificación del oferente", "formulario", form),
        _entry("Anexo 2 Declaración jurada", "declaracion", flat, firma=True),
        _entry("Anexo 5 Oferta Económica", "formulario", offer_form),
        _entry("Anexo 3 Experiencia", "formulario", legacy),
        _entry("Fotocopia de la cédula de identidad", "documento"),
        _entry("Fichas técnicas de los equipos ofertados", "documento"),
        _entry("Boleta de garantía de seriedad", "garantia"),
    ]
    read = _ready_read(dsn, operator, case, annexes, checklist)
    pid = _product(dsn, f"PX{uuid.uuid4().int % 10**4:04d}Q")
    sheet = f"sheet{uuid.uuid4().hex[:12]}"
    from test_v2_quote_case_archive import COT

    drive.add(sheet, "PX224.pdf", COT, mime="application/pdf", data=_minimal_pdf("Ficha PX224"))
    owner_rows(dsn, "insert into catalog.product_document (product_id, kind, title, drive_file_id, mime_type, source) "
                    "values (%s, 'datasheet', 'Ficha PX224', %s, 'application/pdf', 'operator') returning id", (pid, sheet))
    revision = _offer(dsn, operator, case, [("Balanza analítica", 2, 30000, pid), ("Agitador orbital", 1, 40000, None)])
    _bidder(dsn, operator)
    _bidder_document(dsn, store, operator, "cedula", _minimal_pdf(f"Cedula {uuid.uuid4().hex[:6]}"), "cedula.pdf")
    prep = _fill_prep(dsn, operator, case, read, annexes, revision)
    return case, prep, {"form": form, "flat": flat, "offer_form": offer_form, "legacy": legacy}


def test_a_fill_writes_every_output_the_package_and_files_it_in_drive(dsn, store, operator) -> None:
    import zipfile

    drive = TenderDrive()
    case, prep, docs = _full_fill(dsn, store, operator, drive)
    folder = _case_folder(drive, case)
    script = Script()
    with _conn(dsn) as conn:
        preparer = BidPreparer(conn, store, script, MODELS, now=_now, drive=drive, drive_target=_drive_target())
        assert preparer.run(uuid.UUID(prep)) == "ready"
    status, code, result, models, _, _ = _row(dsn, prep)
    outputs = _outputs(dsn, prep)
    kinds = sorted(o["kind"] for o in outputs)
    assert kinds == sorted(["formulario", "formulario_generado", "oferta_economica", "ficha_tecnica",
                            "documento_oferente", "paquete"])
    for o in outputs:
        ext = o["filename"].rsplit(".", 1)[-1].lower()
        assert o["storage_path"] == f"opportunities/{case}/bid/{prep}/{o['file_key']}.{ext}"
        assert re.fullmatch(r"[a-z0-9_-]{1,64}", o["file_key"])
        store.get(o["storage_path"])  # every output is stored
    by_kind = {o["kind"]: o for o in outputs}
    assert by_kind["formulario"]["source_document_id"] == docs["form"]
    assert by_kind["formulario"]["filename"] == "Anexo 1 Identificación.docx"
    assert by_kind["formulario"]["completar_count"] == 0
    assert by_kind["oferta_economica"]["source_document_id"] == docs["offer_form"]
    assert by_kind["formulario_generado"]["source_document_id"] == docs["flat"]
    assert by_kind["paquete"]["filename"] == f"Oferta_{_code(dsn, case)}.zip"
    # The package: every other output, numbered in the checklist's order.
    with zipfile.ZipFile(io.BytesIO(store.get(by_kind["paquete"]["storage_path"]))) as z:
        names = z.namelist()
    assert len(names) == 5 and all(re.match(r"^0[1-5]_", n) for n in names)
    assert names[0] == "01_Anexo 1 Identificación.docx"
    # Checklist estados, and R13 for the legacy form.
    estados = {c["documento"]: c["estado"] for c in result["checklist"]}
    assert estados["Anexo 1 Identificación del oferente"] == "listo"
    assert estados["Anexo 2 Declaración jurada"] == "firmar"
    assert estados["Anexo 5 Oferta Económica"] == "listo"
    assert estados["Anexo 3 Experiencia"] == "falta"
    assert estados["Fotocopia de la cédula de identidad"] == "listo"
    assert estados["Fichas técnicas de los equipos ofertados"] == "listo"
    assert estados["Boleta de garantía de seriedad"] == "falta"
    assert estados["Falta: convertir a Word (.docx): Anexo 3 Experiencia.doc"] == "falta"
    # Three forms filled by the model, each call recorded as step `fill`; the .doc never reached it.
    assert _steps(dsn, prep) == ["fill", "fill", "fill"] and script.calls == 3
    assert models == {"fill": MODELS.fill}
    assert result["offer"]["total"] == 119000 and result["drive"] == "archivado"
    offer_request = next(r for r in script.requests if "Oferta Económica" in r["messages"][0]["content"])
    assert '"total": 119000' in offer_request["messages"][0]["content"]
    # Drive: <case folder>/Licitación/Oferta_<code>.zip, its id on the package output.
    [sub] = [i for i in drive.items.values() if i["parents"] == [folder] and i["name"] == "Licitación"]
    [zipped] = [i for i in drive.items.values() if i["parents"] == [sub["id"]]]
    assert by_kind["paquete"]["drive_file_id"] == zipped["id"] and zipped["mimeType"] == "application/zip"


def _code(dsn, case: str) -> str:
    return owner_rows(dsn, "select n.codigo_externo from crm.opportunity_evidence e join procurement.notice n "
                           "on n.id = e.notice_id where e.opportunity_id = %s", (case,))[0][0]


def test_a_fill_without_drive_skips_it_and_without_an_offer_form_generates_one(dsn, store, operator) -> None:
    import zipfile

    from docx import Document

    case = _tender_case(dsn, operator)
    bases = _annex(dsn, store, operator, case, _minimal_pdf(f"Bases {uuid.uuid4().hex[:6]}"))
    read = _ready_read(dsn, operator, case, [bases], [_entry("Oferta económica", "formulario")])
    revision = _offer(dsn, operator, case, [("Balanza analítica", 2, 30000, None), ("Agitador orbital", 1, 40000, None)])
    _bidder(dsn, operator)
    prep = _fill_prep(dsn, operator, case, read, [bases], revision)
    script = Script()
    with _conn(dsn) as conn:
        assert BidPreparer(conn, store, script, MODELS, now=_now).run(uuid.UUID(prep)) == "ready"
    result, outputs = _row(dsn, prep)[2], _outputs(dsn, prep)
    assert script.calls == 0 and result["drive"] == "apagado"
    offer = next(o for o in outputs if o["kind"] == "oferta_economica")
    assert offer["filename"] == f"Oferta económica {_code(dsn, case)}.docx" and "source_document_id" not in offer
    text = "\n".join(p.text for p in Document(io.BytesIO(store.get(offer["storage_path"]))).paragraphs)
    assert "Oferente Ejemplo" in text and "11.111.111-1" in text
    paquete = next(o for o in outputs if o["kind"] == "paquete")
    assert "drive_file_id" not in paquete
    with zipfile.ZipFile(io.BytesIO(store.get(paquete["storage_path"]))) as z:
        assert z.namelist() == [f"01_Oferta económica {_code(dsn, case)}.docx"]
    assert result["checklist"][0]["estado"] == "listo"


def test_a_fill_stopped_at_the_cap_keeps_what_it_wrote(dsn, store, operator) -> None:
    """Review Focus 2: the offer (no model) is written; the form's call would cross the cap, so it
    never starts; the row ends budget_stopped with the offer kept and no package."""
    month = datetime(2026, 3, 10, 15, 0, tzinfo=UTC)
    owner_rows(dsn, "insert into platform.model_usage (feature, model, step, usd_cost, created_at) "
                    "values ('licitaciones', 'claude-opus-5-5', 'fill', 59.99, %s) returning id", (month,))
    case = _tender_case(dsn, operator)
    form = _annex(dsn, store, operator, case, _docx(f"Nombre del oferente: ________ {uuid.uuid4().hex[:4]}"),
                  filename="Anexo 1.docx", content_type=DOCX, ext="docx")
    read = _ready_read(dsn, operator, case, [form], [_entry("Anexo 1", "formulario", form)])
    revision = _offer(dsn, operator, case, [("Balanza analítica", 1, 100000, None)])
    _bidder(dsn, operator)
    prep = _fill_prep(dsn, operator, case, read, [form], revision)
    script = Script()
    with _conn(dsn) as conn:
        assert BidPreparer(conn, store, script, MODELS, now=lambda: month).run(uuid.UUID(prep)) == "budget_stopped"
    status, code, result, *_ = _row(dsn, prep)
    outputs = _outputs(dsn, prep)
    assert (status, code, script.calls) == ("budget_stopped", "budget_reached", 0)
    assert [o["kind"] for o in outputs] == ["oferta_economica"]
    store.get(outputs[0]["storage_path"])
    assert result["checklist"][0]["estado"] == "falta"


def test_a_fill_on_an_offer_of_another_case_is_refused(dsn, store, operator) -> None:
    case = _tender_case(dsn, operator)
    other = _tender_case(dsn, operator)
    bases = _annex(dsn, store, operator, case, _minimal_pdf(f"Bases {uuid.uuid4().hex[:6]}"))
    read = _ready_read(dsn, operator, case, [bases], [])
    _bidder(dsn, operator)
    prep = _fill_prep(dsn, operator, case, read, [bases], _offer(dsn, operator, other, [("X", 1, 1000, None)]))
    with _conn(dsn) as conn:
        assert BidPreparer(conn, store, Script(), MODELS, now=_now).run(uuid.UUID(prep)) == "failed"
    assert _row(dsn, prep)[:2] == ("failed", "offer_mismatch")


def test_a_package_over_the_bound_fails_with_its_code_and_keeps_the_files(dsn, store, operator, monkeypatch) -> None:
    from origenlab_worker.tender import package

    monkeypatch.setattr(package, "MAX_PACKAGE_BYTES", 10)
    case = _tender_case(dsn, operator)
    bases = _annex(dsn, store, operator, case, _minimal_pdf(f"Bases {uuid.uuid4().hex[:6]}"))
    read = _ready_read(dsn, operator, case, [bases], [])
    _bidder(dsn, operator)
    prep = _fill_prep(dsn, operator, case, read, [bases], _offer(dsn, operator, case, [("X", 1, 1000, None)]))
    with _conn(dsn) as conn:
        assert BidPreparer(conn, store, Script(), MODELS, now=_now).run(uuid.UUID(prep)) == "failed"
    assert _row(dsn, prep)[:2] == ("failed", "package_too_large")
    assert [o["kind"] for o in _outputs(dsn, prep)] == ["oferta_economica"]


def test_the_sweep_claims_fill_rows(dsn, store, operator) -> None:
    case = _tender_case(dsn, operator)
    bases = _annex(dsn, store, operator, case, _minimal_pdf(f"Bases {uuid.uuid4().hex[:6]}"))
    read = _ready_read(dsn, operator, case, [bases], [])
    _bidder(dsn, operator)
    prep = _fill_prep(dsn, operator, case, read, [bases], _offer(dsn, operator, case, [("X", 1, 1000, None)]))
    with _conn(dsn) as conn:
        counts = run_bid_sweep(conn, store, Script(), MODELS, now=_now)
    assert counts["claimed"] >= 1 and _row(dsn, prep)[0] == "ready"


def test_a_fill_call_lost_after_it_was_sent_is_recorded_as_step_fill_and_keeps_the_outputs(dsn, store, operator) -> None:
    from decimal import Decimal

    from origenlab_worker.tender.job import call_worst_case

    case = _tender_case(dsn, operator)
    form = _annex(dsn, store, operator, case, _docx(f"Nombre del oferente: ________ {uuid.uuid4().hex[:4]}"),
                  filename="Anexo 1.docx", content_type=DOCX, ext="docx")
    read = _ready_read(dsn, operator, case, [form], [_entry("Anexo 1", "formulario", form)])
    _bidder(dsn, operator)
    prep = _fill_prep(dsn, operator, case, read, [form], _offer(dsn, operator, case, [("X", 1, 1000, None)]))
    reader = Reader(error=CallInterrupted())
    with _conn(dsn) as conn:
        assert BidPreparer(conn, store, reader, MODELS, now=_now).run(uuid.UUID(prep)) == "failed"
    assert _row(dsn, prep)[:2] == ("failed", "request_failed")
    assert [o["kind"] for o in _outputs(dsn, prep)] == ["oferta_economica"]
    rows = owner_rows(dsn, "select step, output_tokens, usd_cost::text from platform.model_usage "
                           "where bid_preparation_id = %s", (prep,))
    assert {r[0] for r in rows} == {"fill"} and len(rows) == 2 and all(r[1] == 16000 for r in rows)
    assert sum(Decimal(r[2]) for r in rows) == call_worst_case(reader.requests[0], MODELS.fill)


def test_a_failed_match_call_leaves_the_read_ready_and_unmatched(dsn, store, operator) -> None:
    model = f"PX{uuid.uuid4().int % 10**4:04d}W"
    _product(dsn, model)
    case, doc, prep = _queued(dsn, store, operator)
    answer = _answer(doc)
    answer["items"][0]["descripcion"] = f"Balanza {model}"

    def reader(request):
        if "<<<ITEMS>>>" in request["messages"][0]["content"]:
            raise _status_error(529)
        return Reader(answer)(request)

    with _conn(dsn) as conn:
        assert BidPreparer(conn, store, reader, MODELS, now=_now).run(uuid.UUID(prep)) == "ready"
    result = _row(dsn, prep)[2]
    assert result["match"] == {"estado": "fallo", "codigo": "request_failed"}
    assert result["items"][0]["product_id"] is None and len(result["items"][0]["candidatos"]) == 1


def test_a_case_without_a_drive_folder_files_nothing(dsn, store, operator) -> None:
    drive = TenderDrive()
    case = _tender_case(dsn, operator)
    bases = _annex(dsn, store, operator, case, _minimal_pdf(f"Bases {uuid.uuid4().hex[:6]}"))
    read = _ready_read(dsn, operator, case, [bases], [])
    _bidder(dsn, operator)
    prep = _fill_prep(dsn, operator, case, read, [bases], _offer(dsn, operator, case, [("X", 1, 1000, None)]))
    with _conn(dsn) as conn:
        preparer = BidPreparer(conn, store, Script(), MODELS, now=_now, drive=drive, drive_target=_drive_target())
        assert preparer.run(uuid.UUID(prep)) == "ready"
    assert _row(dsn, prep)[2]["drive"] == "sin_carpeta" and drive.writes == []
