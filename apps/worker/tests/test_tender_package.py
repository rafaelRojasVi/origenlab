"""The fill phase's pure pieces (Task 12): the economic offer from a quote revision, each checklist
line's estado, the package's names and zip, the generated offer document. Every value is invented."""

from __future__ import annotations

import io
import zipfile
from datetime import UTC, datetime
from decimal import Decimal

import pytest
from docx import Document

from origenlab_worker.tender import package
from origenlab_worker.tender.package import (
    PackageTooLarge,
    bidder_kinds_for,
    build_package,
    checklist_status,
    is_datasheet_entry,
    is_offer_entry,
    offer_docx,
    offer_payload,
    package_filename,
    package_name,
)


def _revision(**over):
    revision = {
        "quote_number": "01300-26", "approved_at": datetime(2026, 10, 19, 14, 0, tzinfo=UTC),
        "created_at": datetime(2026, 10, 18, 14, 0, tzinfo=UTC), "tax_rate": Decimal("0.19"),
        "price_decimals": 0, "tax_base": Decimal("100000"), "tax_total": Decimal("19000"),
        "grand_total": Decimal("119000"),
        "lines": [
            {"line_kind": "item", "description": "Balanza analítica", "quantity": Decimal("2.000000"),
             "unit_price_qc": Decimal("30000.0000"), "line_total_qc": Decimal("60000.0000")},
            {"line_kind": "logistics", "description": "Flete interno", "quantity": Decimal("1"),
             "unit_price_qc": Decimal("0"), "line_total_qc": Decimal("0")},
            {"line_kind": "item", "description": "Agitador orbital", "quantity": Decimal("1"),
             "unit_price_qc": Decimal("40000"), "line_total_qc": Decimal("40000")},
        ],
    }
    return revision | over


def test_offer_payload_two_lines_neto_iva_total() -> None:
    payload = offer_payload(_revision())
    assert payload == {
        "numero": "01300-26",
        "fecha": "2026-10-19",
        "lineas": [
            {"descripcion": "Balanza analítica", "cantidad": 2, "precio_unitario_neto": 30000, "neto": 60000,
             "iva": 11400, "total": 71400},
            {"descripcion": "Agitador orbital", "cantidad": 1, "precio_unitario_neto": 40000, "neto": 40000,
             "iva": 7600, "total": 47600},
        ],
        "neto": 100000, "iva": 19000, "total": 119000,
    }


def test_offer_payload_rounds_iva_per_line_half_up_and_keeps_a_fractional_quantity() -> None:
    revision = _revision(tax_base=None, tax_total=None, grand_total=None, approved_at=None, lines=[
        {"line_kind": "item", "description": "Reactivo", "quantity": Decimal("2.5"),
         "unit_price_qc": Decimal("1010"), "line_total_qc": Decimal("2525")}])
    payload = offer_payload(revision)
    assert payload["fecha"] == "2026-10-18"
    assert payload["lineas"] == [{"descripcion": "Reactivo", "cantidad": "2.5", "precio_unitario_neto": 1010,
                                  "neto": 2525, "iva": 480, "total": 3005}]
    assert (payload["neto"], payload["iva"], payload["total"]) == (2525, 480, 3005)


def test_checklist_status_listo_falta_firmar() -> None:
    entry = {"documento": "Anexo 1", "tipo": "formulario", "requiere_firma": False}
    done = [{"file_key": "formulario_01", "completar_count": 0}]
    assert checklist_status(entry, done, []) == "listo"
    assert checklist_status(entry, [], []) == "falta"
    assert checklist_status(entry, [{"file_key": "formulario_01", "completar_count": 2}], []) == "falta"
    assert checklist_status(entry | {"requiere_firma": True}, done, []) == "firmar"
    cedula = {"documento": "Fotocopia de cédula de identidad", "tipo": "documento", "requiere_firma": False}
    assert checklist_status(cedula, [], [{"id": "d1", "kind": "cedula"}]) == "listo"


def test_a_missing_document_is_falta_even_when_it_must_be_signed() -> None:
    """Nothing to sign yet: «Firmar» only once there is a document."""
    entry = {"documento": "Declaración jurada", "tipo": "declaracion", "requiere_firma": True}
    assert checklist_status(entry, [], []) == "falta"
    assert checklist_status(entry, [{"file_key": "f", "completar_count": 1}], []) == "falta"


def test_the_offer_entry_is_a_form_named_as_the_economic_offer() -> None:
    assert is_offer_entry({"documento": "Anexo N°5 Oferta Económica", "tipo": "formulario"})
    assert is_offer_entry({"documento": "Formulario económico", "tipo": "formulario"})
    assert is_offer_entry({"documento": "OFERTA ECONOMICA", "tipo": "formulario"})
    assert not is_offer_entry({"documento": "Oferta económica", "tipo": "documento"})
    assert not is_offer_entry({"documento": "Formulario técnico", "tipo": "formulario"})


def test_bidder_documents_answer_the_lines_that_ask_for_them() -> None:
    assert bidder_kinds_for({"documento": "Fotocopia de la cédula de identidad", "tipo": "documento"}) == {"cedula"}
    assert bidder_kinds_for({"documento": "Certificado de inscripción en ChileProveedores",
                             "tipo": "certificado"}) == {"certificado_inscripcion"}
    assert bidder_kinds_for({"documento": "Declaración jurada simple", "tipo": "declaracion"}) == {"declaracion_jurada"}
    assert bidder_kinds_for({"documento": "Experiencia: listado de clientes", "tipo": "documento"}) == {"experiencia"}
    # A line whose form is one of the anexos is filled, not answered with a standard document.
    assert bidder_kinds_for({"documento": "Declaración jurada", "tipo": "declaracion",
                             "archivo_formulario_id": "x"}) == set()
    assert bidder_kinds_for({"documento": "Boleta de garantía", "tipo": "garantia"}) == set()


def test_datasheet_lines() -> None:
    assert is_datasheet_entry({"documento": "Fichas técnicas de los equipos ofertados", "tipo": "documento"})
    assert is_datasheet_entry({"documento": "Catálogo del producto", "tipo": "documento"})
    assert not is_datasheet_entry({"documento": "Ficha técnica", "tipo": "formulario", "archivo_formulario_id": "x"})


def test_package_name_keeps_the_bases_own_name_sanitised() -> None:
    assert package_name("1234567-12-LE26", "Anexo 2 Formulario.docx", 1) == "01_Anexo 2 Formulario.docx"
    assert package_name("1234567-12-LE26", "../a/b\\Anexo:3?.pdf", 12) == "12_Anexo_3_.pdf"
    assert package_name("1234567-12-LE26", "\x00\x01", 3) == "03_documento_1234567-12-LE26"
    assert len(package_name("c", "x" * 400 + ".pdf", 4)) <= 124
    assert package_name("c", "x" * 400 + ".pdf", 4).endswith(".pdf")
    assert package_filename("1234567-12-LE26") == "Oferta_1234567-12-LE26.zip"


def test_build_package_numbers_members_and_refuses_too_large(monkeypatch) -> None:
    data = build_package([("Anexo 1.docx", b"uno"), ("Ficha técnica PX224.pdf", b"%PDF-dos")], "1234567-12-LE26")
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert z.namelist() == ["01_Anexo 1.docx", "02_Ficha técnica PX224.pdf"]
        assert z.read("02_Ficha técnica PX224.pdf") == b"%PDF-dos"
    monkeypatch.setattr(package, "MAX_PACKAGE_BYTES", 5)
    with pytest.raises(PackageTooLarge):
        build_package([("a.pdf", b"123"), ("b.pdf", b"456")], "c")


def test_offer_docx_has_the_bidder_the_number_and_the_lines() -> None:
    payload = offer_payload(_revision())
    data = offer_docx("1234567-12-LE26", payload, {"legal_name": "Oferente Ejemplo", "rut": "11111111-1"})
    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
    assert "Oferta económica 1234567-12-LE26" in text
    assert "OrigenLab" in text and "Oferente Ejemplo" in text and "11.111.111-1" in text and "01300-26" in text
    assert "Balanza analítica" in cells and "$60.000" in cells and "$119.000" in cells
    assert "Flete interno" not in cells
