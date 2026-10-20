"""Finding and filling the buyer's form fields (Task 11). Pure: files are built in the test, every
value is invented."""

from __future__ import annotations

import io
import zipfile

import fitz
import openpyxl
import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from origenlab_worker.tender import forms
from origenlab_worker.tender.forms import (
    COMPLETAR_RE,
    FormError,
    Slot,
    apply_values,
    count_completar,
    form_slots,
    generated_docx,
    inspect_form,
)


def docx_bytes(build) -> bytes:
    doc = Document()
    build(doc)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def xlsx_bytes(build) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    build(ws)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_docx_underscore_paragraph_becomes_a_slot_and_is_filled_and_highlighted() -> None:
    data = docx_bytes(lambda d: (d.add_paragraph("Anexo 1"), d.add_paragraph("Nombre del oferente: ____________")))
    kind, slots = form_slots("anexo.docx", data)
    assert kind == "docx"
    assert [(s.slot_id, s.label) for s in slots] == [("p1", "Nombre del oferente")]
    filled = apply_values("docx", data, {"p1": "Oferente Ejemplo"})
    doc = Document(io.BytesIO(filled))
    assert doc.paragraphs[1].text == "Nombre del oferente: Oferente Ejemplo"
    assert doc.paragraphs[0].text == "Anexo 1"
    assert not any(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in doc.paragraphs[1].runs)
    unknown = Document(io.BytesIO(apply_values("docx", data, {"p1": "[completar: nombre]"})))
    assert any(r.font.highlight_color == WD_COLOR_INDEX.YELLOW for r in unknown.paragraphs[1].runs)


def test_docx_two_blanks_in_one_paragraph_get_two_slots() -> None:
    data = docx_bytes(lambda d: d.add_paragraph("Nombre: ______ RUT: ______"))
    _, slots = form_slots("a.docx", data)
    assert len(slots) == 2 and slots[0].slot_id == "p0" and slots[1].slot_id != "p0"
    filled = apply_values("docx", data, {slots[0].slot_id: "Oferente Ejemplo", slots[1].slot_id: "61.111.111-8"})
    assert Document(io.BytesIO(filled)).paragraphs[0].text == "Nombre: Oferente Ejemplo RUT: 61.111.111-8"


def test_docx_empty_cell_beside_a_label_cell_is_a_slot() -> None:
    def build(d):
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Razón social"
        t.cell(1, 0).text = "RUT"
        t.cell(1, 1).text = "____"

    data = docx_bytes(build)
    _, slots = form_slots("a.docx", data)
    assert [(s.slot_id, s.label) for s in slots] == [("t0r0c1", "Razón social"), ("t0r1c1", "RUT")]
    filled = Document(io.BytesIO(apply_values("docx", data, {"t0r0c1": "Oferente Ejemplo", "t0r1c1": "61.111.111-8"})))
    assert filled.tables[0].cell(0, 1).text == "Oferente Ejemplo"
    assert filled.tables[0].cell(1, 1).text == "61.111.111-8"


def test_xlsx_cell_right_of_a_label_is_the_slot() -> None:
    def build(ws):
        ws["B3"] = "RUT"
        ws["B4"] = "Dirección"
        ws["C4"] = "ya escrito"
        ws["B5"] = "Teléfono"
        ws["C5"] = "ya escrito"

    data = xlsx_bytes(build)
    kind, slots = form_slots("f.xlsx", data)
    assert kind == "xlsx"
    by_id = {s.slot_id: s.label for s in slots}
    assert by_id["Hoja1!C3"] == "RUT"
    assert "Hoja1!C4" not in by_id
    filled = openpyxl.load_workbook(io.BytesIO(apply_values("xlsx", data, {"Hoja1!C3": "61.111.111-8"})))
    cell = filled["Hoja1"]["C3"]
    assert cell.value == "61.111.111-8" and cell.fill.fgColor.rgb in (None, "00000000")


def test_xlsx_below_a_label_when_right_is_taken() -> None:
    def build(ws):
        ws["A1"] = "Descripción"
        ws["B1"] = "Precio"
    _, slots = form_slots("f.xlsx", xlsx_bytes(build))
    assert {s.slot_id for s in slots} == {"Hoja1!A2", "Hoja1!B2"}


def fillable_pdf() -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 60), "Nombre del oferente:")
    w = fitz.Widget()
    w.field_name = "nombre"
    w.field_type = fitz.PDF_WIDGET_TYPE_TEXT
    w.rect = fitz.Rect(200, 45, 400, 65)
    page.add_widget(w)
    return doc.tobytes()


def test_fillable_pdf_round_trips() -> None:
    data = fillable_pdf()
    kind, slots = form_slots("f.pdf", data)
    assert kind == "pdf_fillable" and [s.slot_id for s in slots] == ["nombre"]
    filled = apply_values("pdf_fillable", data, {"nombre": "Oferente Ejemplo"})
    values = {w.field_name: w.field_value for w in fitz.open(stream=filled, filetype="pdf")[0].widgets()}
    assert values == {"nombre": "Oferente Ejemplo"}
    assert count_completar(apply_values("pdf_fillable", data, {"nombre": "[completar: nombre]"}), "pdf_fillable") == 1


def test_flat_pdf_is_not_fillable_with_a_reason() -> None:
    doc = fitz.open()
    doc.new_page().insert_text((50, 60), "Nombre: ________")
    result = inspect_form("flat.pdf", doc.tobytes())
    assert (result.kind, result.slots) == ("not_fillable", [])
    assert result.reason


@pytest.mark.parametrize("name,data", [
    ("bases.doc", b"\xd0\xcf\x11\xe0" + b"0" * 100),
    ("x.docx", b"not a zip"),
    ("x.xlsx", b"PK\x03\x04garbage"),
    ("x.pdf", b"%PDF-broken"),
    ("x.odt", b"abc"),
    ("x.docx", b""),
])
def test_unreadable_or_legacy_forms_are_not_fillable_never_raise(name: str, data: bytes) -> None:
    result = inspect_form(name, data)
    assert result.kind == "not_fillable" and result.slots == [] and result.reason
    assert form_slots(name, data) == ("not_fillable", [])


def test_oversized_file_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(forms, "MAX_FORM_BYTES", 10)
    assert form_slots("a.docx", b"x" * 11)[0] == "not_fillable"


def test_zip_bomb_docx_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    data = docx_bytes(lambda d: d.add_paragraph("Nombre: ____"))
    monkeypatch.setattr(forms, "MAX_UNCOMPRESSED_BYTES", 100)
    assert form_slots("a.docx", data)[0] == "not_fillable"


def test_slot_count_is_capped(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(forms, "MAX_SLOTS", 3)
    data = docx_bytes(lambda d: [d.add_paragraph(f"Campo {i}: ____") for i in range(10)])
    _, slots = form_slots("a.docx", data)
    assert len(slots) == 3


def test_a_docx_without_blanks_is_not_fillable() -> None:
    data = docx_bytes(lambda d: d.add_paragraph("Solo texto"))
    assert form_slots("a.docx", data) == ("not_fillable", [])


def test_apply_values_never_returns_empty_bytes() -> None:
    with pytest.raises(FormError):
        apply_values("not_fillable", b"", {})
    with pytest.raises(FormError):
        apply_values("docx", b"garbage", {"p0": "x"})


def test_generated_docx_starts_with_the_note_and_the_source_name() -> None:
    data = generated_docx("Anexo 1", [("Datos", "Oferente Ejemplo"), ("Garantía", "[completar: monto]")], "anexo1.pdf")
    doc = Document(io.BytesIO(data))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "El original no tiene campos rellenables; este documento reproduce su contenido llenado."
    assert "anexo1.pdf" in texts[1]
    assert "Oferente Ejemplo" in texts and "Datos" in texts
    assert count_completar(data, "docx") == 1


def test_count_completar_counts_docx_and_xlsx() -> None:
    data = docx_bytes(lambda d: (d.add_paragraph("a [completar: x] b"), d.add_table(rows=1, cols=1).cell(0, 0).add_paragraph("[completar: y]")))
    assert count_completar(data, "docx") == 2
    assert count_completar(xlsx_bytes(lambda ws: (ws.__setitem__("A1", "[completar: x]"), ws.__setitem__("A2", "[completar: y]"))), "xlsx") == 2
    assert COMPLETAR_RE.findall("[completar: a] y [completar: b]") == ["[completar: a]", "[completar: b]"]
    assert count_completar(b"junk", "docx") == 0


def test_xlsx_known_value_plain_unknown_highlighted() -> None:
    data = xlsx_bytes(lambda ws: ws.__setitem__("B3", "RUT"))
    out = openpyxl.load_workbook(io.BytesIO(apply_values("xlsx", data, {"Hoja1!C3": "[completar: RUT]"})))
    assert out["Hoja1"]["C3"].fill.fgColor.rgb.endswith("FFFF00")


def test_count_completar_xlsx_with_a_far_cell_is_fast() -> None:
    import time

    def build(ws):
        ws["A1"] = "[completar: x]"
        ws["XFD1048576"] = "z"
    data = xlsx_bytes(build)
    start = time.monotonic()
    assert count_completar(data, "xlsx") == 1
    assert time.monotonic() - start < 5


def test_xlsx_formula_values_are_written_as_text() -> None:
    data = xlsx_bytes(lambda ws: ws.__setitem__("B3", "RUT"))
    out = openpyxl.load_workbook(io.BytesIO(apply_values("xlsx", data, {"Hoja1!C3": '=HYPERLINK("http://x","y")'})))
    cell = out["Hoja1"]["C3"]
    assert cell.data_type == "s" and cell.value == '=HYPERLINK("http://x","y")'


def test_docx_paragraph_with_a_hyperlink_is_left_alone() -> None:
    from docx.oxml import OxmlElement

    def build(d):
        p = d.add_paragraph("Nombre: ____")
        link = OxmlElement("w:hyperlink")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "enlace"
        r.append(t)
        link.append(r)
        p._p.append(link)
        d.add_paragraph("RUT: ____")

    data = docx_bytes(build)
    _, slots = form_slots("a.docx", data)
    assert [s.slot_id for s in slots] == ["p1"]
    out = Document(io.BytesIO(apply_values("docx", data, {"p0": "X", "p1": "61.111.111-8"})))
    assert "enlace" in out.paragraphs[0]._p.xml
    assert out.paragraphs[0].text.startswith("Nombre: ____")
    assert out.paragraphs[1].text == "RUT: 61.111.111-8"


def test_toc_dot_leader_is_not_a_blank() -> None:
    data = docx_bytes(lambda d: d.add_paragraph("Capítulo 1 ........... 12"))
    assert form_slots("a.docx", data) == ("not_fillable", [])


def test_member_caps(monkeypatch: pytest.MonkeyPatch) -> None:
    assert forms.MAX_UNCOMPRESSED_BYTES == 40 * 1024 * 1024
    data = docx_bytes(lambda d: d.add_paragraph("Nombre: ____"))
    monkeypatch.setattr(forms, "MAX_MEMBER_BYTES", 100)
    assert form_slots("a.docx", data)[0] == "not_fillable"


def test_generated_docx_with_control_chars_does_not_raise() -> None:
    data = generated_docx("T\x00", [("H\x00", "a\x00b\x0b")], "x\x00.pdf")
    assert Document(io.BytesIO(data)).paragraphs


def test_pdf_combobox_value_outside_options_becomes_completar() -> None:
    doc = fitz.open()
    page = doc.new_page()
    w = fitz.Widget()
    w.field_name = "region"
    w.field_type = fitz.PDF_WIDGET_TYPE_COMBOBOX
    w.choice_values = ["Norte", "Sur"]
    w.rect = fitz.Rect(100, 100, 300, 120)
    page.add_widget(w)
    data = doc.tobytes()
    out = apply_values("pdf_fillable", data, {"region": "Marte"})
    assert count_completar(out, "pdf_fillable") == 1


def test_count_completar_dedupes_merged_docx_cells() -> None:
    def build(d):
        t = d.add_table(rows=1, cols=2)
        m = t.cell(0, 0).merge(t.cell(0, 1))
        m.text = "[completar: x]"
    assert count_completar(docx_bytes(build), "docx") == 1
